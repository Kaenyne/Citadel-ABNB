"""Build model/ABNB_pitch_model.xlsx: the one-workbook view of the team's ABNB forecasts.

Run from the repo root:   py -3.13 analysis/src/model_workbook/build.py [--no-recalc]
Each tab lives in tabs/tab_<name>.py and exposes build(wb). The workbook is written with openpyxl and then opened
once in Excel (COM) to calculate and cache every formula so that viewers without a calc engine see values.
"""
from __future__ import annotations
import sys, importlib, argparse, time
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE / "tabs"))

import style  # noqa: E402
from openpyxl import Workbook  # noqa: E402

OUT = style.REPO / "model" / "ABNB_pitch_model.xlsx"

TABS = [
    "tab_scenarios",   # built first: the Income Statement lookups need its ranges; moved to the end on save
    "tab_cover",
    "tab_is",
    "tab_revenue",
    "tab_margins",
    "tab_ops",
    "tab_street",
    "tab_stock",
    "tab_rdcf_mgmt",
    "tab_rdcf_market",
]
ORDER = ["Cover", "Income Statement", "Revenue Model", "Margins", "Operating Schedules", "5 Nov & Street",
         "Stock Chart", "Reverse DCF - Mgmt", "Reverse DCF - Market", "Scenario Data"]


def recalc_with_excel(path: Path) -> bool:
    """Open in a fresh hidden Excel instance, full-rebuild calculate, save, and release every COM reference."""
    try:
        import win32com.client  # type: ignore
        import pythoncom  # type: ignore
    except ImportError:
        print("win32com not available; skipping Excel recalculation")
        return False
    import gc
    pythoncom.CoInitialize()
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    xl.AskToUpdateLinks = False
    ok = False
    pid = None
    try:
        import win32process  # type: ignore
        pid = win32process.GetWindowThreadProcessId(xl.Hwnd)[1]
    except Exception:
        pass
    try:
        wb = xl.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=False)
        try:
            xl.CalculateFullRebuild()
            wb.Save()
            ok = True
        finally:
            wb.Close(SaveChanges=False)
            del wb
    finally:
        xl.Quit()
        del xl
        gc.collect()
        pythoncom.CoUninitialize()
        if pid:
            import subprocess, time as _t
            _t.sleep(1)
            subprocess.run(["taskkill", "/PID", str(pid), "/F"], capture_output=True)
    return ok


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-recalc", action="store_true")
    ap.add_argument("--only", nargs="*", help="build only these tab modules (for testing)")
    ap.add_argument("--out", default=str(OUT))
    a = ap.parse_args(argv)
    t0 = time.time()
    wb = Workbook()
    wb.remove(wb.active)
    for name in (a.only or TABS):
        mod = importlib.import_module(name)
        mod.build(wb)
        print(f"  built {name} ({time.time()-t0:.1f}s)")
    # display order
    present = [n for n in ORDER if n in wb.sheetnames] + [n for n in wb.sheetnames if n not in ORDER]
    wb._sheets = [wb[n] for n in present]
    wb.active = 0
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print("saved", out)
    if not a.no_recalc:
        ok = recalc_with_excel(out)
        print("Excel recalculation:", "done" if ok else "skipped")
    return 0


if __name__ == "__main__":
    sys.exit(main())
