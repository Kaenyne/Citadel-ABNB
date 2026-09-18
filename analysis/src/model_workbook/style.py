"""Shared styling helpers for the ABNB pitch workbook (model/ABNB_pitch_model.xlsx).

Conventions (banker standard, stated on the Cover tab):
  blue font   = hard-coded source value (from a repo CSV; the source path is in the row note or the Sources block)
  black font  = formula computed in the sheet
  green font  = link to another sheet
  grey italic = note / label
  light-yellow fill = forecast column;  actual columns unfilled
Units: USD millions unless the row label says otherwise; percentages stored as fractions with a % format.
"""
from __future__ import annotations
from pathlib import Path
import math
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[3]
D = REPO / "data" / "processed"

FONT = "Arial"
NAVY = "1F3864"
BLUE_INPUT = "0000FF"
GREEN_LINK = "007A33"
GREY_NOTE = "6E6E6E"
FILL_HEADER = PatternFill("solid", fgColor=NAVY)
FILL_SECTION = PatternFill("solid", fgColor="D9E1F2")
FILL_FORECAST = PatternFill("solid", fgColor="FFF9E5")
FILL_KEY = PatternFill("solid", fgColor="E2EFDA")
FILL_WARN = PatternFill("solid", fgColor="FCE4D6")
FILL_SHORT = PatternFill("solid", fgColor="F8CBAD")
FILL_STREET = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")
MED = Side(style="medium", color=NAVY)

FMT_M = '#,##0;(#,##0);"-"'
FMT_M1 = '#,##0.0;(#,##0.0);"-"'
FMT_M2 = '#,##0.00;(#,##0.00);"-"'
FMT_PCT = '0.0%;(0.0%);"-"'
FMT_PCT2 = '0.00%;(0.00%);"-"'
FMT_PP = '+0.0"pp";-0.0"pp";"0.0pp"'
FMT_PP2 = '+0.00"pp";-0.00"pp";"0.00pp"'
FMT_EPS = '$0.00;($0.00);"-"'
FMT_USD = '$#,##0.00'
FMT_USD0 = '$#,##0'
FMT_X = '0.0"x"'
FMT_INT = '#,##0'
FMT_DATE = 'dd-mmm-yy'
FMT_PROB = '0.00'

def f_title(): return Font(name=FONT, size=14, bold=True, color=NAVY)
def f_sub(): return Font(name=FONT, size=10, italic=True, color=GREY_NOTE)
def f_hdr(): return Font(name=FONT, size=10, bold=True, color="FFFFFF")
def f_section(): return Font(name=FONT, size=11, bold=True, color=NAVY)
def f_label(bold=False): return Font(name=FONT, size=10, bold=bold)
def f_input(bold=False): return Font(name=FONT, size=10, bold=bold, color=BLUE_INPUT)
def f_formula(bold=False): return Font(name=FONT, size=10, bold=bold, color="000000")
def f_link(bold=False): return Font(name=FONT, size=10, bold=bold, color=GREEN_LINK)
def f_note(): return Font(name=FONT, size=9, italic=True, color=GREY_NOTE)

KIND_FONT = {"input": f_input, "formula": f_formula, "link": f_link, "label": f_label}


def setup(ws, title: str, subtitle: str = "", widths: dict | None = None, label_width: int = 44):
    """Title block in rows 1-2, returns the next free row (4)."""
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = f_title()
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = f_sub()
    ws.column_dimensions["A"].width = label_width
    if widths:
        for c, w in widths.items():
            ws.column_dimensions[c if isinstance(c, str) else get_column_letter(c)].width = w
    return 4


def section(ws, row: int, text: str, ncols: int = 12, note: str = ""):
    """A shaded section bar spanning ncols columns. Returns row+1."""
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = FILL_SECTION
        cell.border = Border(bottom=MED)
    ws.cell(row=row, column=1, value=text).font = f_section()
    if note:
        ws.cell(row=row, column=2, value=note).font = f_note()
    return row + 1


def header(ws, row: int, labels: list, first_col: int = 2, label: str = "", height: float | None = None):
    """Column-header row (navy). labels go in first_col.. Returns row+1."""
    if label:
        c = ws.cell(row=row, column=1, value=label)
        c.font = f_hdr(); c.fill = FILL_HEADER
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=first_col + i, value=lab)
        c.font = f_hdr(); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if height:
        ws.row_dimensions[row].height = height
    return row + 1


def _isnan(v):
    return v is None or (isinstance(v, float) and math.isnan(v))


def write_row(ws, row: int, label: str, values: list, fmt: str = FMT_M, kind: str = "input",
              first_col: int = 2, bold: bool = False, indent: int = 0, forecast_from: int | None = None,
              note: str = "", note_col: int | None = None, fill=None, kinds: list | None = None,
              fmts: list | None = None, wrap_text: bool = False, height: float | None = None):
    """Write a labelled row. values may contain None (skipped), numbers, or strings beginning with '=' (formulas).
    kinds: optional per-value kind list overriding `kind`. forecast_from: index in values from which cells get the
    forecast fill. Returns row+1."""
    lc = ws.cell(row=row, column=1, value=("    " * indent) + label)
    lc.font = f_label(bold=bold)
    for i, v in enumerate(values):
        if _isnan(v):
            continue
        if hasattr(v, "item"):
            v = v.item()
        c = ws.cell(row=row, column=first_col + i, value=v)
        k = kinds[i] if kinds else kind
        if isinstance(v, str) and v.startswith("="):
            k = "formula" if k == "input" else k
        c.font = KIND_FONT.get(k, f_formula)(bold=bold)
        c.number_format = fmts[i] if fmts else fmt
        if isinstance(v, str) and not v.startswith("="):
            c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=bool(wrap_text))
        else:
            c.alignment = Alignment(horizontal="right")
        if forecast_from is not None and i >= forecast_from:
            c.fill = FILL_FORECAST
        if fill is not None:
            c.fill = fill
    if note:
        nc = ws.cell(row=row, column=note_col or (first_col + len(values) + 1), value=safe_text(note))
        nc.font = f_note()
        nc.alignment = Alignment(horizontal="left", vertical="top")
    if height:
        ws.row_dimensions[row].height = height
    return row + 1


def safe_text(s):
    """A string that starts with '=' would be written as a formula and corrupt the file."""
    return (" " + s) if isinstance(s, str) and s.startswith("=") else s


def text_row(ws, row: int, text: str, col: int = 1, font=None, wrap_cols: int = 0, height: float | None = None):
    c = ws.cell(row=row, column=col, value=safe_text(text))
    c.font = font or f_note()
    c.alignment = Alignment(wrap_text=bool(wrap_cols), vertical="top")
    if wrap_cols:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + wrap_cols - 1)
    if height:
        ws.row_dimensions[row].height = height
    return row + 1


def bullets(ws, row: int, lines: list, col: int = 1, wrap_cols: int = 10, font=None, height: float | None = None):
    for ln in lines:
        row = text_row(ws, row, "• " + ln, col=col, font=font or f_label(), wrap_cols=wrap_cols, height=height)
    return row


def table(ws, row: int, df: pd.DataFrame, first_col: int = 1, fmts: dict | None = None, kind: str = "input",
          col_widths: dict | None = None, header_height: float | None = 30,
          default_fmt: str | None = None, wrap_text_cols: list | None = None, fills: dict | None = None,
          row_fill_col: str | None = None, row_fills: dict | None = None):
    """Dump a DataFrame as a formatted table starting at (row, first_col). fmts maps column name -> number format.
    row_fill_col/row_fills: colour whole rows by the value in one column. Returns the next free row."""
    cols = list(df.columns)
    for j, name in enumerate(cols):
        c = ws.cell(row=row, column=first_col + j, value=str(name))
        c.font = f_hdr(); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if header_height:
        ws.row_dimensions[row].height = header_height
    row += 1
    for _, r in df.iterrows():
        rf = None
        if row_fill_col and row_fills:
            rf = row_fills.get(r[row_fill_col])
        for j, name in enumerate(cols):
            v = r[name]
            if _isnan(v):
                continue
            if hasattr(v, "item"):
                v = v.item()
            c = ws.cell(row=row, column=first_col + j, value=v)
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                c.font = f_input() if kind == "input" else f_formula()
                c.number_format = (fmts or {}).get(name, default_fmt or (FMT_M1 if isinstance(v, float) else FMT_INT))
                c.alignment = Alignment(horizontal="right")
            else:
                c.font = f_label()
                if wrap_text_cols and name in wrap_text_cols:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            if fills and name in fills:
                c.fill = fills[name]
            if rf is not None:
                c.fill = rf
        row += 1
    if col_widths:
        for name, w in col_widths.items():
            if name in cols:
                ws.column_dimensions[get_column_letter(first_col + cols.index(name))].width = w
    return row


def freeze(ws, cell="B4"):
    ws.freeze_panes = cell


def col(i: int) -> str:
    return get_column_letter(i)


def sources_block(ws, row: int, items: list, ncols: int = 10):
    """items = [(label, repo path or doc), ...]"""
    row = section(ws, row, "Sources (paths relative to the repo root)", ncols=ncols)
    for lab, path in items:
        ws.cell(row=row, column=1, value=lab).font = f_label()
        ws.cell(row=row, column=2, value=path).font = f_note()
        row += 1
    return row + 1


def pct(x):
    """CSV percentages are stored as 9.89 meaning 9.89%; convert to a fraction for Excel % formats."""
    return None if _isnan(x) else x / 100.0


def read(rel: str, **kw) -> pd.DataFrame:
    return pd.read_csv(D / rel, **kw)
