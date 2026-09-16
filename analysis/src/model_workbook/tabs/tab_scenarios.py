"""'Scenario Data' sheet: the quarterly forecast block for every scenario, keyed '<scenario>|<line>' in column A.
The Income Statement tab pulls its forecast columns from here with INDEX/MATCH driven by a dropdown."""
import style as S
import data as Dt
from openpyxl.styles import Alignment

SHEET = "Scenario Data"
FIRST_DATA_ROW = 5
KEY_COL = "A"
QCOL0 = 4  # column D = 3Q26


def build(wb):
    ws = wb.create_sheet(SHEET)
    S.setup(ws, "Scenario data block (feeds the Income Statement dropdown)",
            "One row per scenario x line item, 3Q26-4Q27. Key in column A = '<scenario key>|<line key>'. Values in USD m unless stated; "
            "percent lines stored as fractions. Sources: 40_lines_quarterly.csv (8 scenarios) and 40_short_case_quarterly.csv (short case; "
            "the 4Q26-marketing-cut variant is derived from 40_short_case_summary.csv q4_marketing_cut_musd).", label_width=48)
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 36
    for i in range(6):
        ws.column_dimensions[S.col(QCOL0 + i)].width = 11
    S.header(ws, 4, ["Scenario", "Line"] + Dt.FC_Q, first_col=2, label="key")
    sq = Dt.scenario_quarterly()
    r = FIRST_DATA_ROW
    for disp, key in Dt.SCENARIOS.items():
        blk = sq[sq.scenario == key].set_index("quarter")
        for csv_col, label in Dt.LINES:
            ws.cell(row=r, column=1, value=f"{key}|{csv_col}").font = S.f_note()
            ws.cell(row=r, column=2, value=disp).font = S.f_label()
            ws.cell(row=r, column=3, value=label.strip()).font = S.f_label()
            for i, q in enumerate(Dt.FC_Q):
                v = blk.loc[q, csv_col] if csv_col in blk.columns else None
                if v is None or v != v:
                    continue
                v = float(v)
                fmt = S.FMT_M1
                if csv_col.endswith("_pct"):
                    v = v / 100.0
                    fmt = S.FMT_PCT2
                elif csv_col in ("eps", "ops_per_booking"):
                    fmt = S.FMT_M2
                c = ws.cell(row=r, column=QCOL0 + i, value=v)
                c.font = S.f_input(); c.number_format = fmt
                c.alignment = Alignment(horizontal="right")
            r += 1
        r += 1  # blank row between scenarios
    last = r
    # scenario list for the dropdown (columns L:M)
    ws.cell(row=4, column=12, value="Scenario names (dropdown list)").font = S.f_hdr()
    ws.cell(row=4, column=12).fill = S.FILL_HEADER
    ws.cell(row=4, column=13, value="key").font = S.f_hdr()
    ws.cell(row=4, column=13).fill = S.FILL_HEADER
    for i, (disp, key) in enumerate(Dt.SCENARIOS.items()):
        ws.cell(row=5 + i, column=12, value=disp).font = S.f_label()
        ws.cell(row=5 + i, column=13, value=key).font = S.f_note()
    ws.column_dimensions["L"].width = 36
    ws.column_dimensions["M"].width = 28
    S.freeze(ws, "D5")
    # remember ranges for other tabs
    wb._scen_ranges = {
        "keys": f"'{SHEET}'!$A${FIRST_DATA_ROW}:$A${last}",
        "data": f"'{SHEET}'!$D${FIRST_DATA_ROW}:$I${last}",
        "qhdr": f"'{SHEET}'!$D$4:$I$4",
        "names": f"'{SHEET}'!$L$5:$L${4 + len(Dt.SCENARIOS)}",
        "namekeys": f"'{SHEET}'!$M$5:$M${4 + len(Dt.SCENARIOS)}",
    }
    return ws
