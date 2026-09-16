"""Recalculate a workbook in Excel and print selected cells (data_only) for checking."""
import sys
sys.path.insert(0, __file__.rsplit("\\", 1)[0] if "\\" in __file__ else ".")
from pathlib import Path
from build import recalc_with_excel
from openpyxl import load_workbook

def main(path, sheet, labels):
    path = Path(path).resolve()
    recalc_with_excel(path)
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    hdr = None
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row):
        a = r[0].value
        if a is None: continue
        if str(a).strip() == "Quarter":
            hdr = [c.value for c in r]
        if any(str(a).strip() == l for l in labels):
            vals = [c.value for c in r]
            print(str(a).strip()[:38].ljust(38), [round(v, 2) if isinstance(v, (int, float)) else v for v in vals[1:]])
    if hdr: print("HDR", hdr[1:])
    errs = 0
    for r in ws.iter_rows():
        for c in r:
            if isinstance(c.value, str) and c.value.startswith("#"):
                errs += 1
                if errs < 10: print("ERR", c.coordinate, c.value)
    print("error cells:", errs)

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3:])
