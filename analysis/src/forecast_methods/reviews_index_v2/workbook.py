"""Build model/ABNB_official_model.xlsx: the nights engine, cell by cell, plus the official income-statement scaffold
with nights as its first line. Every derivable number is an Excel formula; every input is blue with its source beside it.
Palette: Airbnb #FF385C (Rausch) headers, #222222 section bands. Font Arial. Run: python3 workbook.py"""
import json, sys
from pathlib import Path
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.chart import LineChart, BarChart, Reference, Series
sys.path.insert(0, str(Path(__file__).resolve().parent))
import config as C

RAUSCH, INK = "FF385C", "222222"
F_HDR = Font(name="Arial", bold=True, color="FFFFFF", size=10); FILL_HDR = PatternFill("solid", fgColor=RAUSCH)
F_SEC = Font(name="Arial", bold=True, color="FFFFFF", size=10); FILL_SEC = PatternFill("solid", fgColor=INK)
F_IN = Font(name="Arial", color="0000FF", size=10); F_FX = Font(name="Arial", color="000000", size=10)
F_LINK = Font(name="Arial", color="008000", size=10); F_LBL = Font(name="Arial", size=10); F_NOTE = Font(name="Arial", italic=True, color="666666", size=9)
F_BOLD = Font(name="Arial", bold=True, size=10); FILL_FC = PatternFill("solid", fgColor="FFF5F7"); FILL_TODO = PatternFill("solid", fgColor="FFFF00")
FILL_KEY = PatternFill("solid", fgColor="FFE1E7")
PCT, PCT2, NUM1, NUM2, NUM3, MUSD = "0.0%", "0.00%", "0.0", "0.00", "0.000", "#,##0"
OUT = C.OUT; ROOT = C.ROOT; XLSX = ROOT / "model/ABNB_official_model.xlsx"
Q = [f"{q}Q{y}" for y in range(21, 28) for q in range(1, 5)]                     # 1Q21 .. 4Q27 (2021 levels feed the 2022 y/y)
QI = {lab: C.qi(2000 + int(lab[2:]), int(lab[0])) for lab in Q}
COL0 = 3                                                                            # first quarter column = C
col = {lab: COL0 + i for i, lab in enumerate(Q)}
A = lambda lab, r: f"{L(col[lab])}{r}"


def hdr(ws, r, text, cols):
    for c in range(1, cols + 1):
        ws.cell(r, c).fill = FILL_HDR; ws.cell(r, c).font = F_HDR
    ws.cell(r, 1, text)


def sec(ws, r, text, cols):
    for c in range(1, cols + 1):
        ws.cell(r, c).fill = FILL_SEC; ws.cell(r, c).font = F_SEC
    ws.cell(r, 1, text)


def put(ws, r, c, v, font=F_FX, fmt=None, fill=None):
    cell = ws.cell(r, c, v); cell.font = font
    if fmt: cell.number_format = fmt
    if fill: cell.fill = fill
    return cell


def label(ws, r, text, note=None):
    put(ws, r, 1, text, F_LBL)
    if note: put(ws, r, 2, note, F_NOTE)


def build():
    kpi = pd.read_csv(C.KPI); kpi["qi"] = kpi.year.astype(int) * 4 + kpi.q.astype(int) - 1; kpi = kpi.set_index("qi")
    reg = pd.read_csv(OUT / "index_regional_vmatch_pct.csv").set_index("qi")
    w = pd.read_csv(OUT / "mix_weights_stay_quarter.csv").set_index("qi")
    c3 = json.loads((OUT / "stage_c3_3q26.json").read_text()); meta = json.loads((OUT / "final_model_meta.json").read_text())
    stA = pd.read_csv(OUT / "stage_a_tests.csv"); stA = stA[stA.measure == "yoy_vmatch"].set_index("test")
    stB = pd.read_csv(OUT / "stage_b_walkforward.csv"); stB = stB[(stB.measure == C.PRIMARY) & (stB.subset == "full") & (stB.target == "level")].set_index("window")
    kp = pd.read_csv(C.KPI_PANEL).set_index("quarter")
    M = pd.read_csv(C.K2).set_index("northern_stay_quarter")
    guid = pd.read_csv(OUT / "stage_e_guidance.csv")
    regions = ["NAM", "EMEA", "LatAm", "APAC"]

    wb = Workbook(); ws = wb.active; ws.title = "Nights_Engine"
    ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 34
    for lab in Q: ws.column_dimensions[L(col[lab])].width = 9.5
    ncols = COL0 + len(Q) - 1
    hdr(ws, 1, "Airbnb — nights engine (reviews_index_v2, v2.1 stay-quarter mix). Blue = input with source; black = formula; green = link.", ncols)
    put(ws, 2, 1, "Every number below can be traced: inputs carry their file path; every derived cell is an Excel formula on the cells above it. Engine: analysis/src/forecast_methods/reviews_index_v2 (run.py --stage all); note: docs/revenue-forecast-strategy/05_backtests/REVIEWS_INDEX_v2.md", F_NOTE)
    r = 4; sec(ws, r, "Quarter", ncols); put(ws, r, 2, "A = printed, E = engine", F_NOTE)
    for lab in Q:
        put(ws, r, col[lab], lab + ("E" if QI[lab] >= C.qi(2026, 3) else "A"), F_HDR, fill=FILL_HDR).alignment = Alignment(horizontal="center")
    hist_last = C.qi(2026, 2)

    # ---------------- A. printed nights ----------------
    r = 6; sec(ws, r, "A. Printed nights (the KPI, booking-dated)", ncols)
    R_N, R_Y = 7, 8
    label(ws, R_N, "Nights & seats booked (m)", "data/processed/abnb_driver_history_quarterly.csv (10-Q/10-K)")
    label(ws, R_Y, "    y/y", "→ nights / nights four quarters earlier − 1")
    for lab in Q:
        qi = QI[lab]
        if qi <= hist_last and qi in kpi.index: put(ws, R_N, col[lab], float(kpi.loc[qi, "nights_m"]), F_IN, NUM1)
        if qi - 4 in QI.values() and qi <= hist_last and qi in kpi.index:
            prev = [k for k, v in QI.items() if v == qi - 4][0]; put(ws, R_Y, col[lab], f"={A(lab, R_N)}/{A(prev, R_N)}-1", F_FX, PCT2)

    # ---------------- B. the index ----------------
    r = 10; sec(ws, r, "B. Stays index — regional stays growth (same-age two-vintage review counts) × stay-quarter mix weights", ncols)
    R_REG = {}; rr = 11
    for R in regions:
        R_REG[R] = rr; label(ws, rr, f"Stays y/y, {R} (%)", "index_regional_vmatch_pct.csv (Inside Airbnb reviews, vmatch)")
        for lab in Q:
            qi = QI[lab]
            if qi in reg.index and pd.notna(reg.loc[qi, R]): put(ws, rr, col[lab], float(reg.loc[qi, R]), F_IN, NUM2)
        rr += 1
    R_W = {}
    for R in regions:
        R_W[R] = rr; label(ws, rr, f"Stay-mix weight, {R}", "mix_weights_stay_quarter.csv (regional revenue ÷ ADR index; drift rule fwd)")
        for lab in Q:
            qi = QI[lab]
            if qi in w.index: put(ws, rr, col[lab], float(w.loc[qi, R]), F_IN, NUM3)
        rr += 1
    R_IDX = rr; label(ws, R_IDX, "Global stays index, y/y (%)", "→ SUMPRODUCT(regional growth, weights)")
    for lab in Q:
        qi = QI[lab]
        if qi in reg.index and qi in w.index and reg.loc[qi].notna().all():
            c = L(col[lab]); put(ws, R_IDX, col[lab], f"=SUMPRODUCT({c}{R_REG['NAM']}:{c}{R_REG['APAC']},{c}{R_W['NAM']}:{c}{R_W['APAC']})", F_FX, NUM2, FILL_KEY)
    rr = R_IDX + 1

    # ---------------- C. mapping and walk-forward ----------------
    rr += 1; sec(ws, rr, "C. Mapping to the KPI — nights y/y = a + b × index, fitted 1Q23–2Q25 and frozen; scored walk-forward", ncols); rr += 1
    x0, x1 = A("1Q23", R_IDX), A("2Q25", R_IDX); y0, y1 = A("1Q23", R_Y), A("2Q25", R_Y)
    R_A, R_B = rr, rr + 1
    label(ws, R_A, "Intercept a (frozen, 1Q23–2Q25)", "→ INTERCEPT(nights y/y, index) — pre-RNPL quarters only"); put(ws, R_A, COL0, f"=INTERCEPT({y0}:{y1},{x0}:{x1})*100", F_FX, NUM3, FILL_KEY)
    label(ws, R_B, "Slope b (frozen, 1Q23–2Q25)", "→ SLOPE(nights y/y, index)"); put(ws, R_B, COL0, f"=SLOPE({y0}:{y1},{x0}:{x1})*100", F_FX, NUM3, FILL_KEY)
    put(ws, R_A, COL0 + 1, "pp", F_NOTE); put(ws, R_B, COL0 + 1, "pp of nights per index point", F_NOTE)
    rr += 3
    R_SI = rr; label(ws, R_SI, "Stays-implied nights y/y (%)", "→ a + b × index (frozen mapping)")
    R_GAP = rr + 1; label(ws, R_GAP, "Gap: printed − stays-implied (pp)", "the option term I, observed (3Q25 onward)")
    a_ref, b_ref = f"${L(COL0)}${R_A}", f"${L(COL0)}${R_B}"
    for lab in Q:
        qi = QI[lab]
        if qi >= C.qi(2023, 1) and qi <= hist_last:
            put(ws, R_SI, col[lab], f"={a_ref}+{b_ref}*{A(lab, R_IDX)}", F_FX, NUM2)
            put(ws, R_GAP, col[lab], f"={A(lab, R_Y)}*100-{A(lab, R_SI)}", F_FX, NUM2, FILL_KEY if qi >= C.qi(2025, 3) else None)
    rr += 3
    # walk-forward tables
    def wf_block(rr, title, train_start, score_start, score_end):
        sec(ws, rr, title, ncols); rr += 1
        R_P, R_E, R_NE = rr, rr + 1, rr + 2
        label(ws, R_P, "  prediction (refit on quarters before t)", "→ INTERCEPT(prior y, prior x) + SLOPE(prior y, prior x) × index_t")
        label(ws, R_E, "  error: prediction − printed (pp)"); label(ws, R_NE, "  naive error: last quarter's y/y − printed (pp)")
        labs = [q for q in Q if score_start <= QI[q] <= score_end]
        for lab in labs:
            prev = [k for k, v in QI.items() if v == QI[lab] - 1][0]
            ys, xs = f"{A(train_start, R_Y)}:{A(prev, R_Y)}", f"{A(train_start, R_IDX)}:{A(prev, R_IDX)}"
            put(ws, R_P, col[lab], f"=(INTERCEPT({ys},{xs})+SLOPE({ys},{xs})*{A(lab, R_IDX)})*100", F_FX, NUM2)
            put(ws, R_E, col[lab], f"={A(lab, R_P)}-{A(lab, R_Y)}*100", F_FX, NUM2)
            put(ws, R_NE, col[lab], f"=({A(prev, R_Y)}-{A(lab, R_Y)})*100", F_FX, NUM2)
        c0, c1 = A(labs[0], 0)[:-1], A(labs[-1], 0)[:-1]
        R_S = rr + 3
        label(ws, R_S, "  RMSE engine / RMSE naive", f"scored {labs[0]}–{labs[-1]}, n {len(labs)}")
        put(ws, R_S, COL0, f"=SQRT(SUMSQ({c0}{R_E}:{c1}{R_E})/COUNT({c0}{R_E}:{c1}{R_E}))", F_FX, NUM3); put(ws, R_S, COL0 + 1, f"=SQRT(SUMSQ({c0}{R_NE}:{c1}{R_NE})/COUNT({c0}{R_NE}:{c1}{R_NE}))", F_FX, NUM3)
        put(ws, R_S, COL0 + 2, f"={L(COL0)}{R_S}/{L(COL0+1)}{R_S}", F_BOLD, NUM3, FILL_KEY); put(ws, R_S, COL0 + 3, "ratio (pass line ≤ 0.75)", F_NOTE)
        put(ws, R_S + 1, 1, "  mean error (pp)", F_LBL); put(ws, R_S + 1, COL0, f"=AVERAGE({c0}{R_E}:{c1}{R_E})", F_FX, NUM2)
        return rr + 5, R_S, (c0, c1, R_E)
    rr, R_W1, _ = wf_block(rr, "C1. Walk-forward W1 — window 1Q22+, scored 1Q23–2Q26", "1Q22", C.qi(2023, 1), hist_last)
    rr, R_W2, _ = wf_block(rr, "C2. Walk-forward W2 — window 1Q23+, scored 1Q24–2Q26", "1Q23", C.qi(2024, 1), hist_last)
    rr, R_WB, (bc0, bc1, R_BE) = wf_block(rr, "C3. Band — W2 walk-forward on the pre-RNPL quarters only, scored 1Q24–2Q25", "1Q23", C.qi(2024, 1), C.FREEZE_QI)
    R_BAND = R_WB; put(ws, R_BAND, COL0 + 4, "← the RMSE in the first cell of this row is the band (±pp) on every read", F_NOTE)
    band_ref = f"${L(COL0)}${R_BAND}"

    # ---------------- D. 3Q26 read ----------------
    rr += 1; sec(ws, rr, "D. The 3Q26 read — quarter-to-date, day-matched, same-age two vintages, through the frozen mapping", ncols); rr += 1
    R_P3 = {}
    for R in regions:
        R_P3[R] = rr; label(ws, rr, f"  3Q26 QTD stays y/y, {R} (%)", "q3nowcast/E/vintage_matched_nowcast.csv (E6, 3q26_to_date, vmatch_cw)"); put(ws, rr, COL0, float(c3["regional_partial_pct"][R]), F_IN, NUM2)
        put(ws, rr, COL0 + 1, f"={A('3Q26', R_W[R])}", F_LINK, NUM3); put(ws, rr, COL0 + 2, "← 3Q26 stay-mix weight (block B)", F_NOTE); rr += 1
    R_COMP = rr; label(ws, rr, "  Composite QTD index (%)", "→ SUMPRODUCT(regional, weights)/SUM(weights)")
    put(ws, rr, COL0, f"=SUMPRODUCT({L(COL0)}{R_P3['NAM']}:{L(COL0)}{R_P3['APAC']},{L(COL0+1)}{R_P3['NAM']}:{L(COL0+1)}{R_P3['APAC']})/SUM({L(COL0+1)}{R_P3['NAM']}:{L(COL0+1)}{R_P3['APAC']})", F_FX, NUM2); rr += 1
    R_PTF = rr; label(ws, rr, "  Partial-to-full-quarter gap (pp)", "q3nowcast/E/q3_2026_nowcast.csv gap_mean_pp (measured 2023–25)"); put(ws, rr, COL0, float(c3["partial_to_full_gap_pp"]), F_IN, NUM3); rr += 1
    R_FULL = rr; label(ws, rr, "  Full-quarter index (%)", "→ composite + gap"); put(ws, rr, COL0, f"={L(COL0)}{R_COMP}+{L(COL0)}{R_PTF}", F_FX, NUM2); rr += 1
    R_READ = rr; label(ws, rr, "  3Q26 STAYS READ, nights y/y (%)", "→ a + b × full index"); put(ws, rr, COL0, f"={a_ref}+{b_ref}*{L(COL0)}{R_FULL}", F_BOLD, NUM2, FILL_KEY)
    put(ws, rr, COL0 + 1, f"={band_ref}", F_LINK, NUM2); put(ws, rr, COL0 + 2, "± band (block C3)", F_NOTE); rr += 1
    R_LVL = rr; label(ws, rr, "  3Q26 stays level (m)  [lo, hi]", "→ 3Q25 printed × (1 + read)")
    put(ws, rr, COL0, f"={A('3Q25', R_N)}*(1+{L(COL0)}{R_READ}/100)", F_BOLD, NUM1, FILL_KEY)
    put(ws, rr, COL0 + 1, f"={A('3Q25', R_N)}*(1+({L(COL0)}{R_READ}-{band_ref})/100)", F_FX, NUM1); put(ws, rr, COL0 + 2, f"={A('3Q25', R_N)}*(1+({L(COL0)}{R_READ}+{band_ref})/100)", F_FX, NUM1); rr += 1
    read_ref = f"${L(COL0)}${R_READ}"

    # ---------------- E. option term ----------------
    rr += 1; sec(ws, rr, "E. The option term — print = stays + I; I observed 3Q25–2Q26, scenarios for 3Q26, y/y lap and kernel landing", ncols); rr += 1
    R_I0, R_IM, R_IW = rr, rr + 1, rr + 2
    label(ws, R_I0, "  I scenario: 0 (exercise keeps pace with writing)"); put(ws, R_I0, COL0, 0.0, F_IN, NUM2)
    label(ws, R_IM, "  I scenario: mean observed gap", "→ AVERAGE(gaps 3Q25–2Q26)"); put(ws, R_IM, COL0, f"=AVERAGE({A('3Q25', R_GAP)}:{A('2Q26', R_GAP)})", F_FX, NUM2)
    label(ws, R_IW, "  I scenario: largest observed gap (a writing wave)", "→ MAX(gaps)"); put(ws, R_IW, COL0, f"=MAX({A('3Q25', R_GAP)}:{A('2Q26', R_GAP)})", F_FX, NUM2)
    for k, R_ in enumerate([R_I0, R_IM, R_IW]):
        put(ws, R_, COL0 + 1, f"={read_ref}+{L(COL0)}{R_}", F_FX, NUM2); put(ws, R_, COL0 + 2, f"={A('3Q25', R_N)}*(1+{L(COL0+1)}{R_}/100)", F_FX, NUM1)
    put(ws, R_I0, COL0 + 3, "← print y/y and level under each scenario", F_NOTE); rr += 4
    R_YL = rr; label(ws, rr, "  y/y option term = I_t − I_{t−4} (pp)", "what the KPI laps; 3Q26 uses the base print minus the stays read")
    R_BASEP = rr + 1; label(ws, R_BASEP, "  Base print y/y (%) — mechanism (DEC-0029/0019/0025)", "final_nights.md; components in block F")
    for lab in ["3Q26", "4Q26", "1Q27", "2Q27"]:
        prev = f"{lab[0]}Q{int(lab[2:]) - 1}"
        cur = f"({A('3Q26', R_BASEP)}-{read_ref})" if lab == "3Q26" else "0"
        put(ws, R_YL, col[lab], f"={cur}-{A(prev, R_GAP)}", F_FX, NUM2, FILL_KEY)
    rr += 3
    # kernel landing
    label(ws, rr, "  K2 kernel: share of a stay quarter's stays booked 0..3 quarters earlier", "kernel_leadtime_v2/K2_M_matrix.csv"); rr += 1
    R_K = rr
    for i, q in enumerate(["Q1", "Q2", "Q3", "Q4"]):
        put(ws, rr, 1, f"    stays in {q}: lag 0 / 1 / 2 / 3", F_LBL)
        for k in range(4): put(ws, rr, COL0 + k, float(M.loc[q, f"weight_on_GBV_q-{k}"]), F_IN, NUM3)
        rr += 1
    put(ws, rr, 1, "  Landing of each writing wave's cancellations, pp (cohort gap × normalised kernel share by stay quarter)", F_LBL); rr += 1
    waves = [("3Q25", "3Q25"), ("4Q25", "4Q25"), ("1Q26", "1Q26"), ("2Q26", "2Q26")]
    R_LAND0 = rr
    def kshare(bq, k):   # normalised share formula text for a cohort booked in calendar quarter bq (1-4), lag k
        qrow = lambda qn: R_K + (qn - 1)
        terms = [f"{L(COL0+j)}{qrow(((bq - 1 + j) % 4) + 1)}" for j in range(4)]
        return f"{L(COL0+k)}{qrow(((bq - 1 + k) % 4) + 1)}/({'+'.join(terms)})"
    for wl, _ in waves:
        put(ws, rr, 1, f"    written {wl} (gap {wl})", F_LBL); bq = int(wl[0]); i0 = Q.index(wl)
        for k in range(4):
            if i0 + k < len(Q): put(ws, rr, col[Q[i0 + k]], f"=MAX({A(wl, R_GAP)},0)*{kshare(bq, k)}", F_FX, NUM2)
        rr += 1
    put(ws, rr, 1, "    written 3Q26 (base I)", F_LBL); i0 = Q.index("3Q26")
    for k in range(4):
        put(ws, rr, col[Q[i0 + k]], f"=MAX({A('3Q26', R_BASEP)}-{read_ref},0)*{kshare(3, k)}", F_FX, NUM2)
    rr += 1; R_LANDT = rr; label(ws, rr, "    cancellations landing, pp", "column sums")
    for lab in Q[Q.index("3Q25"):Q.index("2Q27") + 1]:
        c = L(col[lab]); put(ws, rr, col[lab], f"=SUM({c}{R_LAND0}:{c}{rr-1})", F_FX, NUM2, FILL_KEY)
    rr += 2

    # ---------------- F. forward path ----------------
    sec(ws, rr, "F. Forward path — base y/y as the sum of sourced components; levels; Street; guidance; the stays path", ncols); rr += 1
    comps = {"3Q26": [("underlying: 0.288 × NA +5.60 + 0.712 × ex-NA +11.62", 9.886, "DEC-0029"), ("fee/cancellation lap", 0.0, ""), ("ex-NA RNPL lap", 0.0, ""), ("events", 0.0, "")],
             "4Q26": [("underlying: team reference", 8.90, "PR #32 / DEC-0019"), ("fee/cancellation lap (ex-NA legs)", -0.78, "DEC-0019"), ("ex-NA RNPL lap", 0.0, ""), ("events", 0.0, "")],
             "1Q27": [("underlying: 0.291 × NA +2.31 + 0.709 × ex-NA +10.773", 8.310, "DEC-0025"), ("fee/cancellation lap", -0.742, "DEC-0025"), ("ex-NA RNPL lap (40% phase-in)", -0.363, "DEC-0025"), ("events", 1.0, "DEC-0025")],
             "2Q27": [("underlying: 0.291 × +2.31 + 0.709 × +10.452", 8.082, "DEC-0025"), ("fee/cancellation lap", -0.742, "DEC-0025"), ("ex-NA RNPL lap (full)", -0.907, "DEC-0025"), ("World Cup lap", -0.5, "DEC-0025")],
             "3Q27": [("underlying: 0.291 × +2.31 + 0.709 × +10.131", 7.855, "DEC-0025"), ("fee/cancellation lap", -0.742, "DEC-0025"), ("ex-NA RNPL lap (full)", -0.907, "DEC-0025"), ("events", 0.0, "")],
             "4Q27": [("underlying: 0.291 × +2.31 + 0.709 × +9.810", 7.627, "DEC-0025"), ("fee/cancellation lap", -0.742, "DEC-0025"), ("ex-NA RNPL lap (full)", -0.907, "DEC-0025"), ("events", 0.0, "")]}
    R_C = [rr, rr + 1, rr + 2, rr + 3]; names = ["  underlying growth (pp)", "  fee / cancellation lap (pp)", "  ex-NA RNPL lap (pp)", "  events (pp)"]
    for i, nm in enumerate(names): label(ws, R_C[i], nm, "components per final_nights.md §4–6")
    for lab, cs in comps.items():
        for i, (desc, v, dec) in enumerate(cs): put(ws, R_C[i], col[lab], v, F_IN, NUM3)
    put(ws, R_C[0], 2, "3Q26: mechanism 0.288×NA 5.60 + 0.712×ex-NA 11.62; 4Q26: ref 8.90; 2027: NA 0.291×2.31 + ex-NA 0.709×g", F_NOTE)
    rr += 4; R_BY = rr; label(ws, rr, "  BASE PRINT y/y (%)", "→ sum of components")
    for lab in comps: c = L(col[lab]); put(ws, rr, col[lab], f"=SUM({c}{R_C[0]}:{c}{R_C[3]})", F_BOLD, NUM3, FILL_KEY)
    # link the base print row in block E to this
    for lab in ["3Q26", "4Q26", "1Q27", "2Q27"]: put(ws, R_BASEP, col[lab], f"={A(lab, R_BY)}", F_LINK, NUM3)
    rr += 1; R_BL = rr; label(ws, rr, "  Base print level (m)", "→ prior-year level × (1 + y/y)")
    for lab in comps:
        prev = f"{lab[0]}Q{int(lab[2:]) - 1}"; prow = R_N if QI[prev] <= hist_last else R_BL
        put(ws, rr, col[lab], f"={A(prev, prow)}*(1+{A(lab, R_BY)}/100)", F_FX, NUM1)
    rr += 1; R_SP = rr; label(ws, rr, "  Stays path y/y (%)", "3Q26 = stays read; 4Q26 on = base (I = 0 at the ceiling)")
    put(ws, rr, col["3Q26"], f"={read_ref}", F_LINK, NUM2)
    for lab in ["4Q26", "1Q27", "2Q27", "3Q27", "4Q27"]: put(ws, rr, col[lab], f"={A(lab, R_BY)}", F_LINK, NUM3)
    rr += 1; R_FY = rr; label(ws, rr, "  FY26 / FY27 nights (m) and FY27 y/y", "FY26 = 1Q26A + 2Q26A + 3Q26E + 4Q26E; FY27 = sum of 2027")
    put(ws, rr, COL0, f"={A('1Q26', R_N)}+{A('2Q26', R_N)}+{A('3Q26', R_BL)}+{A('4Q26', R_BL)}", F_FX, NUM1); put(ws, rr, COL0 + 1, f"={A('1Q27', R_BL)}+{A('2Q27', R_BL)}+{A('3Q27', R_BL)}+{A('4Q27', R_BL)}", F_FX, NUM1)
    put(ws, rr, COL0 + 2, f"={L(COL0+1)}{rr}/{L(COL0)}{rr}-1", F_FX, PCT2); put(ws, rr, COL0 + 3, "FY26 · FY27 · FY27 y/y", F_NOTE)
    rr += 1; R_ST = rr; label(ws, rr, "  Street consensus nights (m)", "Bloomberg MODL 12 Sep 2026 (DEC-0005); 4Q26 per final_nights §4.6")
    put(ws, rr, col["3Q26"], 149.0, F_IN, NUM1); put(ws, rr, col["4Q26"], 134.0, F_IN, NUM1)
    rr += 1; R_STY = rr; label(ws, rr, "  Street y/y (%)", "→ Street level / prior-year printed − 1")
    put(ws, rr, col["3Q26"], f"=({A('3Q26', R_ST)}/{A('3Q25', R_N)}-1)*100", F_FX, NUM2); put(ws, rr, col["4Q26"], f"=({A('4Q26', R_ST)}/{A('4Q25', R_N)}-1)*100", F_FX, NUM2)
    rr += 1; R_G = rr; label(ws, rr, "  Airbnb guidance, nights y/y (low–high, %)", "overnight/02_guidance_ledger.csv (shareholder letters)")
    for _, g in guid[guid.guide_type == "bucket"].iterrows():
        if g.target_period in col: put(ws, rr, col[g.target_period], f"{g.value_low:.0f}–{g.value_high:.0f}", F_IN)
    rr += 1; R_PS = rr; label(ws, rr, "  P(3Q26 print ≥ Street | stays read, band)", "→ 1 − NORMDIST(Street y/y, read, band, TRUE)")
    put(ws, rr, col["3Q26"], f"=1-NORMDIST({A('3Q26', R_STY)},{read_ref},{band_ref},TRUE)", F_BOLD, PCT, FILL_KEY)
    rr += 2

    # ---------------- G. balance sheet ----------------
    sec(ws, rr, "G. RNPL written: unearned fees y/y vs GBV y/y (the option collects no cash until its deadline)", ncols); rr += 1
    R_UF, R_GB, R_SPD, R_Z = rr, rr + 1, rr + 2, rr + 3
    label(ws, R_UF, "  Unearned fees y/y (%)", "overnight/02_kpi_panel_quarterly.csv (10-Q balance sheet)"); label(ws, R_GB, "  GBV y/y, reported (%)", "same file")
    label(ws, R_SPD, "  Spread: unearned fees − GBV (pp)", "→ UF − GBV"); label(ws, R_Z, "  z vs pre-RNPL spread (1Q23–2Q25)", "→ (spread − AVERAGE(pre)) / STDEV(pre)")
    for lab in Q:
        if lab in kp.index and QI[lab] >= C.qi(2023, 1) and QI[lab] <= hist_last:
            put(ws, R_UF, col[lab], float(kp.loc[lab, "unearned_fees_yoy_pct"]), F_IN, NUM2); put(ws, R_GB, col[lab], float(kp.loc[lab, "gbv_yoy_pct"]), F_IN, NUM2)
            put(ws, R_SPD, col[lab], f"={A(lab, R_UF)}-{A(lab, R_GB)}", F_FX, NUM2)
    pre = f"{A('1Q23', R_SPD)}:{A('2Q25', R_SPD)}"
    for lab in ["3Q25", "4Q25", "1Q26", "2Q26"]: put(ws, R_Z, col[lab], f"=({A(lab, R_SPD)}-AVERAGE({pre}))/STDEV({pre})", F_BOLD, NUM2, FILL_KEY)
    put(ws, R_Z, 2, f"pre mean / sd in B: ", F_NOTE); rr += 5

    # ---------------- H. Stage A ----------------
    sec(ws, rr, "H. The instrument is validated where nights are observed — Eurostat panel (computed in the engine; values, sourced)", ncols); rr += 1
    for test, txt in [("A1_elasticity", "β, country fixed effects (se, wild-cluster p, n, clusters)"), ("A2_first_differences", "β in first differences (p)"), ("A5_out_of_sample_by_country", "median walk-forward ratio vs naive (countries ≤ 0.75)")]:
        s = stA.loc[test]; label(ws, rr, "  " + txt, "stage_a_tests.csv"); put(ws, rr, COL0, float(s.value), F_IN, NUM3)
        put(ws, rr, COL0 + 1, (float(s.se_cluster) if pd.notna(s.se_cluster) else None), F_IN, NUM3); put(ws, rr, COL0 + 2, (float(s.p) if pd.notna(s.p) else None), F_IN, NUM3); put(ws, rr, COL0 + 3, int(s.n), F_IN); put(ws, rr, COL0 + 4, str(s.extra), F_NOTE); rr += 1
    rr += 1
    # engine reference values for the checker
    sec(ws, rr, "Engine reference values (Python run) — the checker compares the formulas above to these", ncols); rr += 1
    R_REF = rr
    for nm, v in [("W1 ratio", float(stB.loc["W1", "wf_ratio_vs_naive"])), ("W2 ratio", float(stB.loc["W2", "wf_ratio_vs_naive"])), ("band", float(c3["band_pp"])), ("3Q26 read", float(c3["implied_nights_yoy"])), ("a", float(pd.read_csv(OUT / "stage_c_tests.csv").iloc[0].frozen_a)), ("b", float(pd.read_csv(OUT / "stage_c_tests.csv").iloc[0].frozen_b))]:
        label(ws, rr, "  " + nm, "reviews_index_v2 outputs"); put(ws, rr, COL0, v, F_IN, NUM3); rr += 1

    # ---------------- charts ----------------
    def line_chart(title, rows, labels_, anchor, ymin=None, ymax=None, q_from="1Q23", q_to="4Q27", w_=22, h_=9):
        ch = LineChart(); ch.title = title; ch.height, ch.width = h_, w_; ch.y_axis.title = "%"; ch.style = 2
        c0, c1 = col[q_from], col[q_to]
        for rrow, nm in zip(rows, labels_):
            ser = Series(Reference(ws, min_col=c0, max_col=c1, min_row=rrow), title=nm); ser.smooth = False; ch.series.append(ser)
        ch.set_categories(Reference(ws, min_col=c0, max_col=c1, min_row=4))
        if ymin is not None: ch.y_axis.scaling.min = ymin
        if ymax is not None: ch.y_axis.scaling.max = ymax
        ws.add_chart(ch, anchor)
    # helper rows for charts in percent units: printed y/y ×100 and the combined view row
    rr += 1; R_PY = rr; label(ws, rr, "  chart helper: printed nights y/y (%)", "→ y/y × 100")
    for lab in Q:
        if QI[lab] <= hist_last and QI[lab] >= C.qi(2023, 1): put(ws, rr, col[lab], f"={A(lab, R_Y)}*100", F_FX, NUM2)
    rr += 1; R_VIEW = rr; label(ws, rr, "  chart helper: our view (printed, then base print path)", "")
    for lab in Q:
        if C.qi(2023, 1) <= QI[lab] <= hist_last: put(ws, rr, col[lab], f"={A(lab, R_PY)}", F_FX, NUM2)
        elif lab in comps: put(ws, rr, col[lab], f"={A(lab, R_BY)}", F_FX, NUM2)
    rr += 1; R_STC = rr; label(ws, rr, "  chart helper: Street y/y", "")
    for lab in ["3Q26", "4Q26"]: put(ws, rr, col[lab], f"={A(lab, R_STY)}", F_FX, NUM2)
    rr += 1; R_SPC = rr; label(ws, rr, "  chart helper: stays-implied / stays path", "")
    for lab in Q:
        if C.qi(2023, 1) <= QI[lab] <= hist_last: put(ws, rr, col[lab], f"={A(lab, R_SI)}", F_FX, NUM2)
        elif lab in comps: put(ws, rr, col[lab], f"={A(lab, R_SP)}", F_FX, NUM2)
    anchor_col = L(ncols + 2)
    line_chart("Printed nights y/y vs stays-implied (frozen mapping) — history", [R_PY, R_SPC], ["printed", "stays-implied"], f"{anchor_col}6", q_to="2Q26")
    line_chart("Our view vs the Street — base print path, stays path, Street", [R_VIEW, R_SPC, R_STC], ["our view (print)", "stays path", "Street"], f"{anchor_col}26", ymin=4, ymax=14)
    bc = BarChart(); bc.title = "The option term: printed − stays-implied (pp), 3Q25–2Q26"; bc.height, bc.width = 8, 22
    bc.add_data(Reference(ws, min_col=col["1Q24"], max_col=col["2Q26"], min_row=R_GAP), from_rows=True, titles_from_data=False); bc.set_categories(Reference(ws, min_col=col["1Q24"], max_col=col["2Q26"], min_row=4)); bc.legend = None
    ws.add_chart(bc, f"{anchor_col}46")
    line_chart("RNPL written: unearned fees y/y vs GBV y/y", [R_UF, R_GB], ["unearned fees", "GBV"], f"{anchor_col}64", q_to="2Q26")
    ws.freeze_panes = ws.cell(5, COL0)

    # ---------------- Income statement ----------------
    isws = wb.create_sheet("Income_Statement")
    isws.column_dimensions["A"].width = 52; isws.column_dimensions["B"].width = 30
    per = [q for q in Q if QI[q] >= C.qi(2023, 1)]; fys = ["FY23", "FY24", "FY25", "FY26", "FY27"]
    icol = {p: 3 + i for i, p in enumerate(per)}; fcol = {f: 3 + len(per) + 1 + i for i, f in enumerate(fys)}
    ncol_is = 3 + len(per) + 1 + len(fys)
    hdr(isws, 1, "Airbnb — income statement, official (built line by line; nights first). USD millions unless stated.", ncol_is)
    put(isws, 2, 1, "Row 5 nights: green = link to Nights_Engine. Other lines: labels only until built together (yellow = to build). Quarters A = printed, E = engine forecast.", F_NOTE)
    sec(isws, 4, "Quarter", ncol_is)
    for p in per: put(isws, 4, icol[p], p + ("E" if QI[p] >= C.qi(2026, 3) else "A"), F_HDR, fill=FILL_HDR).alignment = Alignment(horizontal="center")
    for f in fys: put(isws, 4, fcol[f], f, F_HDR, fill=FILL_HDR).alignment = Alignment(horizontal="center")
    for p in per: isws.column_dimensions[L(icol[p])].width = 9.5
    for f in fys: isws.column_dimensions[L(fcol[f])].width = 10
    R_ISN, R_ISY = 5, 6
    label(isws, R_ISN, "Nights & seats booked (m)", "Nights_Engine: printed (A), base print path (E)")
    label(isws, R_ISY, "    y/y", "→ nights / four quarters earlier − 1")
    for p in per:
        src_row = R_N if QI[p] <= hist_last else R_BL
        put(isws, R_ISN, icol[p], f"=Nights_Engine!{A(p, src_row)}", F_LINK, NUM1, FILL_FC if QI[p] > hist_last else None)
        prev = f"{p[0]}Q{int(p[2:]) - 1}"
        if prev in icol: put(isws, R_ISY, icol[p], f"={L(icol[p])}{R_ISN}/{L(icol[prev])}{R_ISN}-1", F_FX, PCT2)
        elif prev in Q: put(isws, R_ISY, icol[p], f"=Nights_Engine!{A(p, R_Y)}", F_LINK, PCT2)
    for f in fys:
        yy = f[2:]; qs = [f"{q}Q{yy}" for q in range(1, 5)]
        put(isws, R_ISN, fcol[f], "=" + "+".join(f"{L(icol[q])}{R_ISN}" for q in qs), F_FX, NUM1, FILL_FC if f in ("FY26", "FY27") else None)
        if f != "FY23": put(isws, R_ISY, fcol[f], f"={L(fcol[f])}{R_ISN}/{L(fcol[fys[fys.index(f)-1]])}{R_ISN}-1", F_FX, PCT2)
    lines = ["Gross booking value ($bn)", "    y/y", "ADR ($ = GBV / nights)", "    y/y", "Revenue", "    y/y", "    Take rate (revenue / GBV)", "",
             "    Cost of revenue (cash)", "    Operations & support (cash)", "    Product development (cash)", "    Sales & marketing (cash)", "    General & administrative (cash, ex lodging-tax reserves)",
             "Total cash costs", "    % of revenue", "    Depreciation & amortisation", "Adjusted EBITDA", "    Adjusted EBITDA margin", "", "    Stock-based compensation", "Operating income (GAAP)",
             "    Interest income", "    Interest expense", "    Other income / (expense), net", "Pre-tax income", "    Income tax", "Net income", "    Diluted shares (m)", "Diluted EPS ($)"]
    rr2 = 8; sec(isws, 7, "Lines to build together (in order) — empty until each has its rationale file", ncol_is)
    for ln in lines:
        if ln: put(isws, rr2, 1, ln, F_BOLD if not ln.startswith("    ") else F_LBL)
        for p in per:
            if ln and QI[p] > hist_last: isws.cell(rr2, icol[p]).fill = FILL_TODO
        rr2 += 1
    isws.freeze_panes = "C5"

    # ---------------- Cover + Sources ----------------
    cv = wb.create_sheet("Cover", 0); cv.column_dimensions["A"].width = 30; cv.column_dimensions["B"].width = 110
    hdr(cv, 1, "Airbnb (ABNB) — official model, built line by line. Line 1: nights.", 2)
    rows_ = [("What this is", "The nights engine (reviews_index_v2, v2.1) laid out cell by cell, and the income-statement scaffold with nights as its first line. Everything derivable is an Excel formula; every input carries its source."),
             ("How to trace", "Nights_Engine, blocks A–H: printed nights → regional stays growth × stay-mix weights = index → INTERCEPT/SLOPE on 1Q23–2Q25 = the frozen mapping → walk-forward W1/W2 reproduced row by row → 3Q26 read → the option term (scenarios, y/y lap, kernel landing) → forward path as sums of sourced components → Street, guidance, tail probability → unearned fees vs GBV → Stage A statistics."),
             ("Colours", "Header fill FF385C = section; fill 222222 = block title; blue text = input with source; black = formula; green = link to another sheet or block; pink fill = a headline number; yellow fill = to be built."),
             ("Engine", "analysis/src/forecast_methods/reviews_index_v2 — python3 run.py --stage all (≈15 s), then python3 workbook.py rebuilds this file. Spec, pre-registration and results: docs/revenue-forecast-strategy/05_backtests/REVIEWS_INDEX_v2.md; rationale: REVIEWS_INDEX_v2_RATIONALE.md."),
             ("Status", "Nights: built (DEC-0029 base 146.8m; engine stays read in Nights_Engine block D). All other income-statement lines: to build together, one at a time, each with its rationale file (DEC-0021 / DEC-0026)."),
             ("Verification", "Formulas were recalculated in Excel and checked against the engine's Python outputs (block 'Engine reference values'): W1/W2 ratios, band, a, b, 3Q26 read.")]
    for i, (k, v) in enumerate(rows_, start=3):
        put(cv, i, 1, k, F_BOLD); put(cv, i, 2, v, F_LBL).alignment = Alignment(wrap_text=True, vertical="top"); cv.row_dimensions[i].height = 48
    src = wb.create_sheet("Sources"); src.column_dimensions["A"].width = 46; src.column_dimensions["B"].width = 90; src.column_dimensions["C"].width = 22
    hdr(src, 1, "Sources and decisions", 3)
    for i, (a_, b_, c_) in enumerate([("Printed nights", "data/processed/abnb_driver_history_quarterly.csv", "DEC-0003"), ("Review counts, two vintages", "data/processed/q3nowcast/E/market_vintage_monthly.csv; q3nowcast_v2/E", "E note; WPK-A"),
        ("Regional stays growth (vmatch)", "data/processed/forecast_methods/reviews_index_v2/index_regional_vmatch_pct.csv", "REVIEWS_INDEX_v2 §2"), ("Stay-mix weights", "reviews_index_v2/mix_weights_stay_quarter.csv ← overnight/10_regional_panel_quarterly.csv", "§5 (v2.1)"),
        ("3Q26 quarter-to-date by region", "data/processed/q3nowcast/E/vintage_matched_nowcast.csv; q3_2026_nowcast.csv", "E6"), ("Lead-time kernel", "data/processed/forecast_methods/kernel_leadtime_v2/K2_M_matrix.csv", "kernel lane"),
        ("Base path components", "docs/pitch-model-v2/lines/final_nights.md §4–6", "DEC-0019 / 0025 / 0029"), ("Street nights", "Bloomberg MODL 12 Sep 2026", "DEC-0005"), ("Guidance", "data/processed/overnight/02_guidance_ledger.csv", "letters"),
        ("Unearned fees, GBV", "data/processed/overnight/02_kpi_panel_quarterly.csv", "RNPL audit"), ("Eurostat panel statistics", "reviews_index_v2/stage_a_tests.csv", "§2.2")], start=2):
        put(src, i, 1, a_, F_LBL); put(src, i, 2, b_, F_LBL); put(src, i, 3, c_, F_LBL)
    wb.save(XLSX); return XLSX, dict(R_W1=R_W1, R_W2=R_W2, R_BAND=R_BAND, R_READ=R_READ, R_A=R_A, R_B=R_B, R_REF=R_REF, COL0=COL0, R_PS=R_PS, col3q26=col["3Q26"])


if __name__ == "__main__":
    path, refs = build(); print("wrote", path); (OUT / "workbook_refs.json").write_text(json.dumps(refs))
