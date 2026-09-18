"""QA gate for the built workbook. Every message is a failure."""
from __future__ import annotations
import argparse, json, re, sys
from pathlib import Path
from openpyxl import load_workbook
from . import spec as specmod

ERR = ("#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NUM!", "#NULL!")
NAME_RE = re.compile(r"\b([A-Z][A-Z0-9_]*_[0-9A-Z]+)\b")
BUILTIN = {"INDEX", "MATCH", "SUM", "MIN", "MAX", "IF", "ABS", "EXP", "LN"}
LICENSED = re.compile(r"(BEST_|BDH\(|BDP\(|=BDS|LSEG Workspace export|Third Bridge)", re.I)
ROW_RE = re.compile(r"^\|\s*(\w+)\s*\|\s*([0-9A-Z]+)\s*\|\s*([-+]?\d*\.?\d+)")

def _dossier_points(path: Path) -> dict[tuple[str, str], float]:
    pts = {}
    if not path.exists():
        return pts
    in_sec = False
    for ln in path.read_text().splitlines():
        if ln.startswith("## 2."):
            in_sec = True; continue
        if in_sec and ln.startswith("## "):
            break
        if in_sec:
            m = ROW_RE.match(ln)
            if m and m.group(1) not in ("scenario", "---"):
                pts[(m.group(1), m.group(2))] = float(m.group(3))
    return pts

def check(workbook_path: str | Path, spec_path: str | Path, root: str | Path | None = None, allow_uncalculated: bool = False) -> list[str]:
    root = Path(root or Path.cwd()); msgs: list[str] = []
    s = specmod.load(spec_path, check_paths=False)
    wbf = load_workbook(workbook_path); wbv = load_workbook(workbook_path, data_only=True)
    names = set(wbf.defined_names.keys())
    # 1. error cells and calculation state
    uncalculated = 0
    for ws in wbf.worksheets:
        wsv = wbv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.strip() in ERR:
                    msgs.append(f"{ws.title}!{c.coordinate}: error literal {v.strip()}")
                if isinstance(v, str) and v.startswith("="):
                    vv = wsv[c.coordinate].value
                    if vv is None:
                        uncalculated += 1
                    elif isinstance(vv, str) and vv.strip() in ERR:
                        msgs.append(f"{ws.title}!{c.coordinate}: evaluates to {vv.strip()}")
                    for n in NAME_RE.findall(v):
                        if n not in names and n.split("_")[0] not in BUILTIN:
                            msgs.append(f"{ws.title}!{c.coordinate}: unresolved name {n}")
                if isinstance(v, str) and LICENSED.search(v):
                    msgs.append(f"{ws.title}!{c.coordinate}: looks like licensed export content: {v[:60]!r}")
    if uncalculated and not allow_uncalculated:
        msgs.append(f"{uncalculated} formula cells have no cached value: run recalc")
    # 2. every spec input: name, provenance files, receipt exit 0, value present in dossier §2
    for lid in s.order:
        ln = s.lines[lid]
        for p in ln.periods:
            if f"{lid}_{p}" not in names:
                msgs.append(f"{lid}_{p}: defined name missing")
        if ln.kind != "input":
            continue
        pv = ln.provenance
        dpath, rpath = root / pv["dossier"], root / pv["receipt"]
        if not dpath.exists():
            msgs.append(f"{lid}: dossier missing {pv['dossier']}"); continue
        if not rpath.exists():
            msgs.append(f"{lid}: receipt missing {pv['receipt']}")
        else:
            rec = json.loads(rpath.read_text())
            if rec.get("exit_code") != 0:
                msgs.append(f"{lid}: receipt exit_code {rec.get('exit_code')}")
        pts = _dossier_points(dpath); tol = float(pv.get("tolerance", 0))
        for sc, per_vals in ln.values.items():
            for p, v in per_vals.items():
                if (sc, p) not in pts:
                    msgs.append(f"{lid} {sc} {p}: value {v} not stated in dossier §2")
                elif abs(pts[(sc, p)] - float(v)) > tol:
                    msgs.append(f"{lid} {sc} {p}: spec {v} vs dossier {pts[(sc, p)]} exceeds tolerance {tol}")
    return msgs

def main(argv=None):
    ap = argparse.ArgumentParser(); ap.add_argument("--wb", required=True); ap.add_argument("--spec", required=True)
    ap.add_argument("--allow-uncalculated", action="store_true")
    a = ap.parse_args(argv)
    msgs = check(a.wb, a.spec, allow_uncalculated=a.allow_uncalculated)
    print("\n".join(msgs) if msgs else "QA clean"); return 1 if msgs else 0

if __name__ == "__main__":
    sys.exit(main())
