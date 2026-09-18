"""Build the pitch model workbook from the spec. Every input is a named cell fed by the
Scenario Data tab through the Scenario selector; every other cell is a formula in names."""
from __future__ import annotations
import argparse, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from . import spec as specmod
from . import decisions as decmod

TABS = {"drivers": "Drivers", "revenue": "Revenue and Guide", "costs": "Costs and Earnings",
        "valuation": "Valuation and Call", "event": "5 Nov Event Card", "street": "Street"}
ORDER = ["Cover", "Drivers", "Revenue and Guide", "Costs and Earnings", "Valuation and Call",
         "5 Nov Event Card", "Street", "Evidence", "Decision Log", "Scenario Data"]
FMT = {"pct": "0.00%", "musd": "#,##0", "usd": "0.00", "m": "0.0", "x": "0.0x", "prob": "0.00", "usd_share": "0.00"}
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # yellow: an input, named, fed by Scenario Data
FORMULA_FILL = PatternFill("solid", fgColor="FFFFFF")
HEAD = Font(bold=True)

def _name(wb: Workbook, name: str, sheet: str, col: int, row: int) -> None:
    wb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet}'!${get_column_letter(col)}${row}")

def build(spec_path: str | Path, out_path: str | Path, decisions_path: str | Path | None = None,
          recalc: bool = False, check_paths: bool = True) -> Path:
    s = specmod.load(spec_path, check_paths=check_paths)
    periods, scenarios = list(s.meta["periods"]), list(s.meta["scenarios"])
    wb = Workbook(); wb.remove(wb.active)
    cover = wb.create_sheet("Cover")
    cover["A1"] = "Airbnb (ABNB) — pitch model v2"; cover["A1"].font = Font(bold=True, size=14)
    cover["A2"] = f"Price date {s.meta['price_date']} · spot ${s.meta['spot']}"
    cover["A4"] = "Scenario"; cover["A4"].font = HEAD; cover["B4"] = scenarios[0]; cover["B4"].fill = INPUT_FILL
    dv = DataValidation(type="list", formula1='"' + ",".join(scenarios) + '"', allow_blank=False)
    cover.add_data_validation(dv); dv.add("B4"); _name(wb, "Scenario", "Cover", 2, 4)
    cover["A6"] = "Yellow cells are inputs; each carries a defined name (id_period), a grade, and a decision id. Every other number is a formula in those names: use Formulas > Trace Precedents."
    # Scenario Data
    sd = wb.create_sheet("Scenario Data")
    for j, h in enumerate(["id", "period", "label", "unit", *scenarios], start=1):
        sd.cell(row=1, column=j, value=h).font = HEAD
    sd_row = {}
    r = 2
    for lid in s.order:
        ln = s.lines[lid]
        if ln.kind != "input":
            continue
        for p in ln.periods:
            sd.cell(row=r, column=1, value=lid); sd.cell(row=r, column=2, value=p)
            sd.cell(row=r, column=3, value=ln.label); sd.cell(row=r, column=4, value=ln.unit)
            for k, sc in enumerate(scenarios):
                sd.cell(row=r, column=5 + k, value=float(ln.values[sc][p]))
            sd_row[(lid, p)] = r; r += 1
    first_sc, last_sc = get_column_letter(5), get_column_letter(4 + len(scenarios))
    # block tabs
    sheets = {b: wb.create_sheet(t) for b, t in TABS.items()}
    rows = {b: 2 for b in TABS}
    for b, ws in sheets.items():
        for j, h in enumerate(["id", "line", "unit", "grade · decision", *periods], start=1):
            ws.cell(row=1, column=j, value=h).font = HEAD
        ws.freeze_panes = "E2"; ws.column_dimensions["B"].width = 44; ws.column_dimensions["D"].width = 18
    for lid in s.order:
        ln = s.lines[lid]
        if ln.block not in sheets:
            raise specmod.SpecError(f"{lid}: unknown block {ln.block}")
        ws = sheets[ln.block]; row = rows[ln.block]; rows[ln.block] += 1
        ws.cell(row=row, column=1, value=lid); ws.cell(row=row, column=2, value=ln.label); ws.cell(row=row, column=3, value=ln.unit)
        if ln.kind == "input":
            pv = ln.provenance; ws.cell(row=row, column=4, value=f"{pv['grade']} · {pv['decision']}")
        else:
            ws.cell(row=row, column=4, value="formula: " + ln.expr)
        for p in ln.periods:
            col = 5 + periods.index(p); c = ws.cell(row=row, column=col)
            if ln.kind == "input":
                sr = sd_row[(lid, p)]
                c.value = f"=INDEX('Scenario Data'!${first_sc}${sr}:${last_sc}${sr},MATCH(Scenario,'Scenario Data'!${first_sc}$1:${last_sc}$1,0))"
                c.fill = INPUT_FILL
            else:
                c.value = "=" + specmod.to_excel(ln.expr, p, periods)
            c.number_format = FMT.get(ln.unit, "General")
            _name(wb, f"{lid}_{p}", ws.title, col, row)
    # Evidence
    ev = wb.create_sheet("Evidence")
    for j, h in enumerate(["id", "line", "kind", "grade", "decision", "tolerance", "dossier", "receipt", "expr"], start=1):
        ev.cell(row=1, column=j, value=h).font = HEAD
    for i, lid in enumerate(s.order, start=2):
        ln = s.lines[lid]; pv = ln.provenance
        vals = [lid, ln.label, ln.kind, pv.get("grade", ""), pv.get("decision", ""), pv.get("tolerance", ""),
                pv.get("dossier", ""), pv.get("receipt", ""), ln.expr]
        for j, v in enumerate(vals, start=1):
            ev.cell(row=i, column=j, value=v)
    # Decision Log
    dl = wb.create_sheet("Decision Log")
    for j, h in enumerate(["id", "line", "period", "scenario", "value", "reason", "rejected", "date"], start=1):
        dl.cell(row=1, column=j, value=h).font = HEAD
    if decisions_path and Path(decisions_path).exists():
        for i, d in enumerate(decmod.load(decisions_path), start=2):
            for j, v in enumerate([d.id, d.line, d.period, d.scenario, d.value, d.reason, d.rejected, d.date], start=1):
                dl.cell(row=i, column=j, value=v)
    wb._sheets = [wb[n] for n in ORDER if n in wb.sheetnames]
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    out = Path(out_path); out.parent.mkdir(parents=True, exist_ok=True); wb.save(out)
    if recalc:
        from . import recalc as rc
        rc.recalc(out)
    return out

def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--spec", default="model/pitch_model_v2/spec/lines.yaml")
    ap.add_argument("--out", default="model/pitch_model_v2/ABNB_pitch_model_v2.xlsx")
    ap.add_argument("--decisions", default="docs/pitch-model-v2/DECISIONS.md")
    ap.add_argument("--no-recalc", action="store_true")
    a = ap.parse_args(argv)
    out = build(a.spec, a.out, a.decisions, recalc=not a.no_recalc)
    print("built", out); return 0

if __name__ == "__main__":
    sys.exit(main())
