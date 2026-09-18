from pathlib import Path
from openpyxl import load_workbook
from pitch_model_v2 import build

FIX = Path(__file__).parent / "fixtures" / "spec_small.yaml"

def test_build_names_and_formulas(tmp_path):
    out = build.build(FIX, tmp_path / "m.xlsx", check_paths=False)
    wb = load_workbook(out)
    names = set(wb.defined_names.keys())
    assert {"D1_3Q26", "D6_3Q26", "R2_4Q26", "R5FY_FY26", "Scenario"} <= names
    ws = wb["Drivers"]
    # find the D6 / 3Q26 cell through its defined name
    dest = wb.defined_names["D6_3Q26"].attr_text  # "'Drivers'!$G$5" style
    sheet, ref = dest.split("!"); cell = wb[sheet.strip("'")][ref.replace("$", "")]
    assert cell.value == "=D1_3Q26*D4T_3Q26"
    r2 = wb.defined_names["R2_3Q26"].attr_text.split("!")
    assert wb[r2[0].strip("'")][r2[1].replace("$", "")].value == "=LAM_3Q26*(2/3*D6_2Q26+1/3*D6_1Q26)"
    d1 = wb.defined_names["D1_3Q26"].attr_text.split("!")
    v = wb[d1[0].strip("'")][d1[1].replace("$", "")].value
    assert v.startswith("=INDEX('Scenario Data'!") and "MATCH(Scenario," in v
    sd = wb["Scenario Data"]
    assert sd["A1"].value == "id" and sd["E1"].value == "base" and sd["F1"].value == "short"
    ev = wb["Evidence"]
    ids = [ev.cell(row=r, column=1).value for r in range(2, ev.max_row + 1)]
    assert "D1" in ids and "D6" in ids
    assert wb.calculation.fullCalcOnLoad is True
