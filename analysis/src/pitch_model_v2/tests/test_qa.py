import json
from pathlib import Path
from openpyxl import load_workbook
from pitch_model_v2 import build, qa

FIX = Path(__file__).parent / "fixtures" / "spec_small.yaml"

def _dossiers(root: Path):
    for lid, val, per in [("D1", "146.3", "3Q26"), ("D4", "176.9", "3Q26"), ("R1", "0.1724", "3Q26")]:
        d = root / "docs/pitch-model-v2/dossiers"; d.mkdir(parents=True, exist_ok=True)
        rd = root / f"data/processed/pitch_model_v2/receipts/{lid}"; rd.mkdir(parents=True, exist_ok=True)
        (rd / "receipt.json").write_text(json.dumps({"id": lid, "exit_code": 0}))
    (root / "docs/pitch-model-v2/dossiers/D1_nights.md").write_text("## 2. The number\n| scenario | period | point |\n|---|---|---|\n| base | 3Q26 | 146.3 |\n| base | 4Q26 | 133.2 |\n| short | 3Q26 | 145.0 |\n| short | 4Q26 | 131.0 |\n| base | 1Q26 | 143.1 |\n| base | 2Q26 | 134.4 |\n| short | 1Q26 | 143.1 |\n| short | 2Q26 | 134.4 |\n")
    (root / "docs/pitch-model-v2/dossiers/D4_adr.md").write_text("## 2. The number\n| scenario | period | point |\n|---|---|---|\n| base | 1Q26 | 176.0 |\n| base | 2Q26 | 179.4 |\n| base | 3Q26 | 176.9 |\n| base | 4Q26 | 171.0 |\n| short | 1Q26 | 176.0 |\n| short | 2Q26 | 179.4 |\n| short | 3Q26 | 174.6 |\n| short | 4Q26 | 169.0 |\n")
    (root / "docs/pitch-model-v2/dossiers/R1_kernel.md").write_text("## 2. The number\n| scenario | period | point |\n|---|---|---|\n| base | 3Q26 | 0.1724 |\n| base | 4Q26 | 0.1203 |\n| short | 3Q26 | 0.1724 |\n| short | 4Q26 | 0.1203 |\n")

def test_qa_passes_on_clean_build(tmp_path):
    _dossiers(tmp_path)
    out = build.build(FIX, tmp_path / "m.xlsx", check_paths=False)
    assert qa.check(out, FIX, root=tmp_path, allow_uncalculated=True) == []

def test_qa_flags_value_outside_tolerance(tmp_path):
    _dossiers(tmp_path)
    (tmp_path / "docs/pitch-model-v2/dossiers/D1_nights.md").write_text("## 2. The number\n| scenario | period | point |\n|---|---|---|\n| base | 3Q26 | 140.0 |\n")
    out = build.build(FIX, tmp_path / "m.xlsx", check_paths=False)
    msgs = qa.check(out, FIX, root=tmp_path, allow_uncalculated=True)
    assert any("D1" in m and "3Q26" in m for m in msgs)

def test_qa_flags_value_not_stated_in_dossier(tmp_path):
    _dossiers(tmp_path)
    (tmp_path / "docs/pitch-model-v2/dossiers/D1_nights.md").write_text("## 2. The number\n| scenario | period | point |\n|---|---|---|\n| base | 3Q26 | 146.3 |\n| base | 4Q26 | 133.2 |\n| short | 3Q26 | 145.0 |\n| base | 1Q26 | 143.1 |\n| base | 2Q26 | 134.4 |\n| short | 1Q26 | 143.1 |\n| short | 2Q26 | 134.4 |\n")
    out = build.build(FIX, tmp_path / "m.xlsx", check_paths=False)
    msgs = qa.check(out, FIX, root=tmp_path, allow_uncalculated=True)
    assert any("not stated in dossier" in m for m in msgs)

def test_qa_flags_error_cells(tmp_path):
    _dossiers(tmp_path)
    out = build.build(FIX, tmp_path / "m.xlsx", check_paths=False)
    wb = load_workbook(out); wb["Drivers"]["Z2"] = "#REF!"; wb.save(out)
    assert any("#REF!" in m for m in qa.check(out, FIX, root=tmp_path, allow_uncalculated=True))

def test_qa_flags_unresolved_name(tmp_path):
    _dossiers(tmp_path)
    out = build.build(FIX, tmp_path / "m.xlsx", check_paths=False)
    wb = load_workbook(out); wb["Drivers"]["Z3"] = "=NOPE_3Q26*2"; wb.save(out)
    assert any("NOPE_3Q26" in m for m in qa.check(out, FIX, root=tmp_path, allow_uncalculated=True))

def test_qa_allows_exp_ln_functions(tmp_path):
    _dossiers(tmp_path)
    out = build.build(FIX, tmp_path / "m.xlsx", check_paths=False)
    wb = load_workbook(out); wb["Drivers"]["Z4"] = "=EXP(D1_3Q26)"; wb.save(out)
    msgs = qa.check(out, FIX, root=tmp_path, allow_uncalculated=True)
    assert not any("unresolved name" in m for m in msgs)


NIGHTS_SPEC = """
meta:
  price_date: "2026-09-16"
  spot: 167.51
  scenarios: [base]
  periods: [1Q26, 2Q26, 3Q26]
lines:
  - id: NIGHTS
    label: Nights (m)
    unit: m
    block: drivers
    kind: input
    periods: [1Q26, 2Q26, 3Q26]
    values:
      base:  {1Q26: 121.1, 2Q26: 115.1, 3Q26: 146.3}
    provenance:
      - dossier: docs/pitch-model-v2/dossiers/H0_history.md
        receipt: data/processed/pitch_model_v2/receipts/H0/receipt.json
        grade: A
        decision: DEC-0001
        tolerance: 0.05
        periods: [1Q26, 2Q26]
        item: nights_m
        scenario_map: {base: actual}
      - dossier: docs/pitch-model-v2/dossiers/D1_nights_3q26.md
        receipt: data/processed/pitch_model_v2/receipts/D1/receipt.json
        grade: B
        decision: DEC-0002
        tolerance: 0.05
        periods: [3Q26]
        item: nights_m
"""

def _nights_dossiers(root: Path, hist_1q26: float = 121.1):
    d = root / "docs/pitch-model-v2/dossiers"; d.mkdir(parents=True, exist_ok=True)
    for lid in ("H0", "D1"):
        rd = root / f"data/processed/pitch_model_v2/receipts/{lid}"; rd.mkdir(parents=True, exist_ok=True)
        (rd / "receipt.json").write_text(json.dumps({"id": lid, "exit_code": 0}))
    (d / "H0_history.md").write_text(
        "# H0 — history\n\n## 2. The number\n(narrative)\n\n"
        "### 2a. Model inputs (machine-readable)\n\n"
        "| item | scenario | period | point | unit | note |\n"
        "|---|---|---|---|---|---|\n"
        f"| nights_m | actual | 1Q26 | {hist_1q26} | m | letter/10-Q KPI box |\n"
        "| nights_m | actual | 2Q26 | 115.1 | m | letter/10-Q KPI box |\n"
        "| adr_usd | actual | 1Q26 | 168.43 | $ | letter/10-Q KPI box |\n"
        "| adr_usd | actual | 2Q26 | 166.01 | $ | letter/10-Q KPI box |\n"
        "\n## 3. Derivation chain\nn/a\n"
    )
    (d / "D1_nights_3q26.md").write_text(
        "# D1 — 3Q26 nights\n\n## 2. The number\n(narrative)\n\n"
        "### 2a. Model inputs (machine-readable)\n\n"
        "| item | scenario | period | point | unit | note |\n"
        "|---|---|---|---|---|---|\n"
        "| nights_m | base | 3Q26 | 146.3 | m nights | bias-corrected read |\n"
        "| nights_m | short | 3Q26 | 145.0 | m nights | print lands low |\n"
        "\n## 3. Derivation chain\nn/a\n"
    )

def test_qa_ties_via_2a_with_item_and_scenario_map(tmp_path):
    _nights_dossiers(tmp_path)
    spec_path = tmp_path / "nights.yaml"; spec_path.write_text(NIGHTS_SPEC)
    out = build.build(spec_path, tmp_path / "n.xlsx", check_paths=False)
    assert qa.check(out, spec_path, root=tmp_path, allow_uncalculated=True) == []
    # Move the history 1Q26 value outside tolerance: exactly one message, naming NIGHTS and 1Q26.
    _nights_dossiers(tmp_path, hist_1q26=999.0)
    msgs = qa.check(out, spec_path, root=tmp_path, allow_uncalculated=True)
    assert len(msgs) == 1
    assert "NIGHTS" in msgs[0] and "1Q26" in msgs[0]

def test_qa_2a_period_all_matches_any_period(tmp_path):
    d = tmp_path / "docs/pitch-model-v2/dossiers"; d.mkdir(parents=True, exist_ok=True)
    rd = tmp_path / "data/processed/pitch_model_v2/receipts/RX"; rd.mkdir(parents=True, exist_ok=True)
    (rd / "receipt.json").write_text(json.dumps({"id": "RX", "exit_code": 0}))
    (d / "RX_rate.md").write_text(
        "# RX — constant rate\n\n## 2. The number\n(narrative)\n\n"
        "### 2a. Model inputs (machine-readable)\n\n"
        "| item | scenario | period | point | unit | note |\n"
        "|---|---|---|---|---|---|\n"
        "| rate_x | base | all | 0.05 | pct | constant across periods |\n"
        "\n## 3. Derivation chain\nn/a\n"
    )
    spec = """
meta:
  price_date: "2026-09-16"
  spot: 167.51
  scenarios: [base]
  periods: [1Q26, 2Q26]
lines:
  - id: RATE
    label: Rate
    unit: pct
    block: drivers
    kind: input
    periods: [1Q26, 2Q26]
    values:
      base: {1Q26: 0.05, 2Q26: 0.05}
    provenance:
      - dossier: docs/pitch-model-v2/dossiers/RX_rate.md
        receipt: data/processed/pitch_model_v2/receipts/RX/receipt.json
        grade: A
        decision: DEC-0001
        tolerance: 0.0001
        item: rate_x
"""
    spec_path = tmp_path / "rate.yaml"; spec_path.write_text(spec)
    out = build.build(spec_path, tmp_path / "r.xlsx", check_paths=False)
    assert qa.check(out, spec_path, root=tmp_path, allow_uncalculated=True) == []
