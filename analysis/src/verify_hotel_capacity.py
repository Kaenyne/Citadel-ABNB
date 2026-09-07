"""Verify two official spreadsheet anchors, retaining room-night units explicitly.

Reads the already downloaded Madrid and Buenos Aires workbooks. No network call.
"""
import csv
import hashlib
import json
from pathlib import Path

from model_hotel_funnel import available_daily_rooms
from research_integrity import finite_number

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / 'data/raw/hotel_funnel_audit/public'
OUT = ROOT / 'data/processed/hotel_funnel_audit'


def hotel_category_sum(total, hotel_categories, excluded_categories):
    values = [total, *hotel_categories, *excluded_categories]
    for value in values:
        finite_number(value, 'Category count', minimum=0)
        if int(value) != value:
            raise ValueError('Category count must be integral')
    if sum(hotel_categories) + sum(excluded_categories) != total:
        raise ValueError('Hotel and excluded categories do not reconcile to total')
    return sum(hotel_categories)


def main():
    from openpyxl import load_workbook
    config = json.loads((ROOT / 'analysis/config/hotel_funnel_audit.json').read_text(encoding='utf-8'))
    anchors = {m['market']: m for m in config['market_anchors']}
    madrid_path = RAW / 'madrid_rooms_2026.xlsx'
    madrid_book = load_workbook(madrid_path, read_only=True, data_only=True)
    madrid_sheet = madrid_book['P1110426']
    july = [row for row in madrid_sheet.values if row[1] == 'Julio']
    if len(july) != 1:
        raise ValueError('Expected exactly one July row in Madrid table')
    m = july[0]
    hotel_rooms = hotel_category_sum(m[2], list(m[3:7]), list(m[7:9]))
    if hotel_rooms != anchors['madrid']['hotel_rooms']:
        raise ValueError('Madrid source differs from configured anchor')
    buenos_path = RAW / 'buenos_rooms.xlsx'
    buenos_book = load_workbook(buenos_path, read_only=True, data_only=True)
    sheet = buenos_book['2026']
    if sheet['B2'].value != 'MARZO' or sheet['H5'].value != 'Boutique':
        raise ValueError('Buenos Aires period or category changed')
    if 'multiplicado' not in str(sheet['A11'].value):
        raise ValueError('Room-night unit definition missing')
    properties, available_nights = sheet['H6'].value, sheet['H7'].value
    if properties != anchors['buenos-aires']['boutique_establishments'] or available_nights != anchors['buenos-aires']['boutique_available_room_nights']:
        raise ValueError('Buenos Aires source differs from configured anchor')
    if sum(sheet.cell(6, col).value for col in range(3, 9)) != sheet['B6'].value:
        raise ValueError('Buenos Aires establishments do not reconcile')
    if sum(sheet.cell(7, col).value for col in range(3, 9)) != sheet['B7'].value:
        raise ValueError('Buenos Aires available room nights do not reconcile')
    rows = [
        dict(market='madrid',source_id='HOTEL-MAD',period='2026-07',metric='hotel_rooms_estimated_open',
             value=hotel_rooms,unit='rooms',calculation='sum hotel star columns D:G; total C minus hostal H:I',
             file=madrid_path.name,sha256=hashlib.sha256(madrid_path.read_bytes()).hexdigest()),
        dict(market='buenos-aires',source_id='HOTEL-BUE',period='2026-03',metric='boutique_establishments',
             value=properties,unit='establishments',calculation='sheet 2026 H6',
             file=buenos_path.name,sha256=hashlib.sha256(buenos_path.read_bytes()).hexdigest()),
        dict(market='buenos-aires',source_id='HOTEL-BUE',period='2026-03',metric='boutique_available_room_nights',
             value=available_nights,unit='available_room_nights_in_month',calculation='sheet 2026 H7; rooms multiplied by days open',
             file=buenos_path.name,sha256=hashlib.sha256(buenos_path.read_bytes()).hexdigest()),
        dict(market='buenos-aires',source_id='HOTEL-BUE',period='2026-03',metric='boutique_average_daily_available_rooms',
             value=available_daily_rooms(available_nights,31),unit='average_daily_available_rooms_not_physical_keys',
             calculation='H7 / 31 calendar days; not a physical room census',
             file=buenos_path.name,sha256=hashlib.sha256(buenos_path.read_bytes()).hexdigest())]
    with (OUT / 'verified_spreadsheet_anchors.csv').open('w', encoding='utf-8', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    madrid_book.close()
    buenos_book.close()
    print('Verified Madrid category totals and Buenos Aires establishments/available room nights; 4 extracted rows')


if __name__ == '__main__':
    main()
