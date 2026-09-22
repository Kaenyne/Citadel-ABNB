"""
41_cost_leg: is there a cost leg to the short? Street-implied cost growth vs history, the cost-surprise record at every print,
where our cost excess sits against LSEG's cost-of-goods consensus, and FY27 cost-growth scenarios.

Run:  py -3.13 analysis/src/margin_build/41_cost_leg/run.py        (exit 0; writes data/processed/margin_build/41_cost_leg/)

Pre-registration: docs/margin-build/notes/41_42_prereg.md (committed f6d01166 before this script existed).
Cost measure everywhere: revenue - adjusted EBITDA ("EBITDA costs"). It is the only cost object the Street publishes implicitly,
it needs no D&A assumption, and for our builds it equals cash costs less D&A. Street = LSEG means (vendor and dates in the inputs).
Reads:  02_financial_panel (actuals), 03_consensus_pit (at-print and current consensus, COGS field), 06_fy27_path_v2 (quarterly 2027
        consensus), 40_line_build (base and evidence_only), model/ABNB_official_model_complete.xlsx (official v2 revenue, cached values).
Writes: 41_annual_cost_growth.csv, 41_quarterly_cost_growth.csv, 41_cost_surprise_history.csv, 41_tests.csv, 41_cogs_attribution.csv,
        41_fy27_scenarios.csv, 41_2h26_scenarios.csv, 41_live_3q26.csv.
"""
from __future__ import annotations
from itertools import combinations
from pathlib import Path
import numpy as np
import pandas as pd
import statsmodels.api as sm
from scipy import stats

ROOT = Path(__file__).resolve().parents[4]
MB = ROOT / "data/processed/margin_build"
OUT = MB / "41_cost_leg"; OUT.mkdir(parents=True, exist_ok=True)
PAN = MB / "02_financial_panel/02_panel_quarterly.csv"
SURP = MB / "03_consensus_pit/03_surprise_history.csv"
CAD = MB / "03_consensus_pit/03_consensus_at_dates.csv"
CUR = MB / "03_consensus_pit/03_current_consensus.csv"
CQ = MB / "06_fy27_path_v2/06_consensus_quarterly_2027.csv"
LB = MB / "40_line_build/40_lines_quarterly.csv"
LBA = MB / "40_line_build/40_annual.csv"
OFFICIAL = ROOT / "model/ABNB_official_model_complete.xlsx"

W1 = ("2023Q1", "2026Q2"); W2 = ("2024Q1", "2026Q2"); RECENT = ("2025Q3", "2026Q2")
ETR_FY27 = 17.5           # M7 sheet, as in 40_line_build
q_to_panel = lambda q: f"{q[-1]}Q{q[2:4]}"          # 2025Q3 -> 3Q25
panel_to_q = lambda q: f"20{q[2:]}Q{q[0]}"          # 3Q25 -> 2025Q3

pan = pd.read_csv(PAN).set_index("quarter")
pan["ebitda_costs"] = pan.revenue - pan.adj_ebitda_reported
cur = pd.read_csv(CUR); cur = cur[cur.vendor == "LSEG"].set_index("period")
cq = pd.read_csv(CQ).set_index("quarter")
lb = pd.read_csv(LB); lba = pd.read_csv(LBA)

def official_v2() -> pd.DataFrame:
    """Official income statement v2 (Theo's pitch model, DEC-0042): revenue and adjusted EBITDA per forecast quarter, cached values."""
    from openpyxl import load_workbook
    ws = load_workbook(OFFICIAL, read_only=True, data_only=True)["Income_Statement"]
    rows = {r: [ws.cell(r, c).value for c in range(1, 30)] for r in range(1, 40)}
    head = rows[4]; col = {h: i for i, h in enumerate(head) if isinstance(h, str) and h.endswith("E")}
    lab = {str(v[0]).strip(): r for r, v in rows.items() if v[0]}
    r_rev = next(r for k, r in lab.items() if k.startswith("Revenue")); r_e = next(r for k, r in lab.items() if k.startswith("Adjusted EBITDA") and "margin" not in k.lower())
    return pd.DataFrame({q[:4]: dict(revenue=float(rows[r_rev][i]), adj_ebitda=float(rows[r_e][i])) for q, i in col.items()}).T

ov2 = official_v2()

# ----------------------------------------------------------------------------------------------------------------------
# 1. Annual cost growth: actuals, Street, our builds
# ----------------------------------------------------------------------------------------------------------------------
def fy_actual(y):
    qs = [f"{i}Q{y}" for i in range(1, 5)]; s = pan.loc[qs]
    return float(s.revenue.sum()), float(s.adj_ebitda_reported.sum())
h1 = pan.loc[["1Q26", "2Q26"]]; h1_rev, h1_e = float(h1.revenue.sum()), float(h1.adj_ebitda_reported.sum())
lbq = lb[lb.scenario == "base"].set_index("quarter"); evq = lb[lb.scenario == "evidence_only"].set_index("quarter")
Q27 = ["1Q27", "2Q27", "3Q27", "4Q27"]
rows = []
for y in ["22", "23", "24", "25"]:
    r, e = fy_actual(y); rows.append(dict(period=f"FY{y}", source="actual", revenue=r, adj_ebitda=e))
for p in ["FY26", "FY27", "FY28"]:
    rows.append(dict(period=p, source="street_annual", revenue=float(cur.loc[p, "revenue_mean"]), adj_ebitda=float(cur.loc[p, "ebitda_mean"]),
                     n=int(cur.loc[p, "ebitda_n"]), asof=cur.loc[p, "as_of_row_date"]))
rows.append(dict(period="FY26", source="street_quarterly_sum", revenue=h1_rev + float(cq.loc[["3Q26", "4Q26"], "revenue_mean_musd"].sum()),
                 adj_ebitda=h1_e + float(cq.loc[["3Q26", "4Q26"], "ebitda_mean_musd"].sum())))
rows.append(dict(period="FY27", source="street_quarterly_sum", revenue=float(cq.loc[Q27, "revenue_mean_musd"].sum()), adj_ebitda=float(cq.loc[Q27, "ebitda_mean_musd"].sum())))
for src, d in [("line_build_base", lbq), ("line_build_evidence_only", evq)]:
    rows.append(dict(period="FY26", source=src, revenue=h1_rev + float(d.loc[["3Q26", "4Q26"], "revenue"].sum()), adj_ebitda=h1_e + float(d.loc[["3Q26", "4Q26"], "adj_ebitda"].sum())))
    rows.append(dict(period="FY27", source=src, revenue=float(d.loc[Q27, "revenue"].sum()), adj_ebitda=float(d.loc[Q27, "adj_ebitda"].sum())))
rows.append(dict(period="FY26", source="official_v2", revenue=h1_rev + float(ov2.loc[["3Q26", "4Q26"], "revenue"].sum()), adj_ebitda=h1_e + float(ov2.loc[["3Q26", "4Q26"], "adj_ebitda"].sum())))
rows.append(dict(period="FY27", source="official_v2", revenue=float(ov2.loc[Q27, "revenue"].sum()), adj_ebitda=float(ov2.loc[Q27, "adj_ebitda"].sum())))
ann = pd.DataFrame(rows)
ann["ebitda_costs"] = ann.revenue - ann.adj_ebitda; ann["margin_pct"] = ann.adj_ebitda / ann.revenue * 100
def prior(r):
    y = int(r.period[2:]) - 1; p = f"FY{y}"
    if r.source == "actual" or (r.period == "FY26"):
        base = ann[(ann.period == p) & (ann.source == "actual")]
    elif r.source.startswith("street"):
        base = ann[(ann.period == p) & (ann.source == ("street_annual" if r.source == "street_annual" else "street_quarterly_sum"))]
    else:
        base = ann[(ann.period == p) & (ann.source == r.source)]
    return base.iloc[0] if len(base) else None
for c in ["revenue_growth_pct", "cost_growth_pct", "incremental_margin_pct"]:
    ann[c] = np.nan
for i, r in ann.iterrows():
    b = prior(r)
    if b is not None:
        ann.loc[i, "revenue_growth_pct"] = (r.revenue / b.revenue - 1) * 100
        ann.loc[i, "cost_growth_pct"] = (r.ebitda_costs / b.ebitda_costs - 1) * 100
        ann.loc[i, "incremental_margin_pct"] = (r.adj_ebitda - b.adj_ebitda) / (r.revenue - b.revenue) * 100
ann.to_csv(OUT / "41_annual_cost_growth.csv", index=False)

# quarterly y/y cost growth: actuals 1Q24-2Q26, Street and line build 3Q26-4Q27
qrows = []
for q in ["1Q24", "2Q24", "3Q24", "4Q24", "1Q25", "2Q25", "3Q25", "4Q25", "1Q26", "2Q26"]:
    p = f"{q[0]}Q{int(q[2:]) - 1}"
    qrows.append(dict(quarter=q, source="actual", revenue_yoy_pct=(pan.loc[q, "revenue"] / pan.loc[p, "revenue"] - 1) * 100,
                      cost_yoy_pct=(pan.loc[q, "ebitda_costs"] / pan.loc[p, "ebitda_costs"] - 1) * 100, margin_pct=pan.loc[q, "adj_ebitda_margin_pct"]))
street_q = {q: (float(cq.loc[q, "revenue_mean_musd"]), float(cq.loc[q, "ebitda_mean_musd"])) for q in ["3Q26", "4Q26"] + Q27}
def q_costs(src, q):
    if q in pan.index: return float(pan.loc[q, "revenue"]), float(pan.loc[q, "ebitda_costs"])
    if src == "street": r, e = street_q[q]; return r, r - e
    d = lbq if src == "line_build_base" else ov2
    r = float(d.loc[q, "revenue"]); e = float(d.loc[q, "adj_ebitda"]); return r, r - e
for src in ["street", "line_build_base", "official_v2"]:
    for q in ["3Q26", "4Q26"] + Q27:
        p = f"{q[0]}Q{int(q[2:]) - 1}"; r, c = q_costs(src, q); rp, cp = q_costs(src, p)
        qrows.append(dict(quarter=q, source=src, revenue_yoy_pct=(r / rp - 1) * 100, cost_yoy_pct=(c / cp - 1) * 100, margin_pct=(r - c) / r * 100))
pd.DataFrame(qrows).to_csv(OUT / "41_quarterly_cost_growth.csv", index=False)

# ----------------------------------------------------------------------------------------------------------------------
# 2. Cost-surprise record at every print (point-in-time LSEG consensus, morning of print)
# ----------------------------------------------------------------------------------------------------------------------
s = pd.read_csv(SURP)
s = s[s.print_quarter >= "2021Q1"].copy()
s["street_costs"] = s.cons_revenue_at_print_musd - s.cons_ebitda_at_print_musd
s["actual_costs"] = s.actual_revenue_musd - s.actual_adj_ebitda_musd
s["revenue_surprise"] = s.actual_revenue_musd - s.cons_revenue_at_print_musd
s["ebitda_surprise"] = s.actual_adj_ebitda_musd - s.cons_ebitda_at_print_musd
s["cost_surprise"] = s.actual_costs - s.street_costs
s["cost_surprise_pct"] = s.cost_surprise / s.street_costs * 100
s["cor_share"] = [float(pan.loc[q_to_panel(q), "cor_cash"] / pan.loc[q_to_panel(q), "revenue"]) for q in s.print_quarter]
s["cost_surprise_exvar"] = s.cost_surprise - s.cor_share * s.revenue_surprise
s["cost_surprise_exvar_pct"] = s.cost_surprise_exvar / s.street_costs * 100
cad = pd.read_csv(CAD); cog = cad[cad.target_role == "printed_q_at_print"].set_index("printed_quarter")
s["street_cogs"] = [float(cog.loc[q, "cogs_mean"]) if q in cog.index else np.nan for q in s.print_quarter]
s["street_cogs_obs_date"] = [cog.loc[q, "cogs_obs_date"] if q in cog.index else None for q in s.print_quarter]
s["actual_cor"] = [float(pan.loc[q_to_panel(q), "cor_cash"]) for q in s.print_quarter]
s["cor_surprise"] = s.actual_cor - s.street_cogs; s["cor_surprise_pct"] = s.cor_surprise / s.street_cogs * 100
s["opex_surprise"] = s.cost_surprise - s.cor_surprise       # everything in costs that is not cost of revenue
keep = ["print_quarter", "print_date", "cons_ebitda_obs_date", "cons_revenue_at_print_musd", "cons_ebitda_at_print_musd", "actual_revenue_musd",
        "actual_adj_ebitda_musd", "revenue_surprise", "ebitda_surprise", "street_costs", "actual_costs", "cost_surprise", "cost_surprise_pct",
        "cor_share", "cost_surprise_exvar", "cost_surprise_exvar_pct", "street_cogs", "street_cogs_obs_date", "actual_cor", "cor_surprise",
        "cor_surprise_pct", "opex_surprise"]
s[keep].to_csv(OUT / "41_cost_surprise_history.csv", index=False)

def window(df, w): return df[(df.print_quarter >= w[0]) & (df.print_quarter <= w[1])]
tests = []
for wname, w in [("W1", W1), ("W2", W2), ("all_2021Q1+", ("2021Q1", "2026Q2"))]:
    d = window(s, w).reset_index(drop=True)
    for var in ["cost_surprise_pct", "cost_surprise_exvar_pct"]:
        x = d[var].to_numpy(); rec = (d.print_quarter >= RECENT[0]).to_numpy(); k = int(rec.sum())
        obs = x[rec].mean() - x[~rec].mean()
        diffs = [x[list(c)].mean() - np.delete(x, list(c)).mean() for c in combinations(range(len(x)), k)]
        p = float(np.mean(np.array(diffs) >= obs - 1e-12))
        tests.append(dict(test="T1_regime_recent4_vs_earlier", window=wname, variable=var, n=len(x), n_recent=k, recent_mean=x[rec].mean(),
                          earlier_mean=x[~rec].mean(), stat=obs, p_one_sided=p, pass_line="p<=0.10 both W1,W2 (raw)", note="split chosen after seeing data"))
        y, xl = x[1:], x[:-1]; ols = sm.OLS(y, sm.add_constant(xl)).fit(cov_type="HC1")
        slope, t = float(ols.params[1]), float(ols.tvalues[1]); p1 = float(stats.t.sf(t, df=len(y) - 2))
        pred_next = float(ols.params[0] + ols.params[1] * x[-1])        # AR(1) forecast for the next print (3Q26) from the last observed surprise
        tests.append(dict(test="T2_persistence_AR1", window=wname, variable=var, n=len(y), stat=slope, t=t, p_one_sided=p1, intercept=float(ols.params[0]),
                          last_value=float(x[-1]), pred_next_print=pred_next, pass_line="slope>0, p<=0.10 both W1,W2"))
        tests.append(dict(test="count_costs_above_street", window=wname, variable=var, n=len(x), stat=int((x > 0).sum()), recent_mean=x[rec].mean(), earlier_mean=x[~rec].mean()))
    c = d.cor_surprise_pct.dropna().to_numpy(); kpos = int((c > 0).sum())
    tests.append(dict(test="T3_cor_above_cogs_consensus", window=wname, variable="cor_surprise_pct", n=len(c), stat=kpos, recent_mean=float(np.mean(c[-4:])),
                      earlier_mean=float(np.mean(c[:-4])), mean_all=float(c.mean()), p_one_sided=float(stats.binomtest(kpos, len(c), 0.5, alternative="greater").pvalue),
                      pass_line="sign test p<=0.10 both W1,W2"))
tests = pd.DataFrame(tests)
for tname in ["T1_regime_recent4_vs_earlier", "T2_persistence_AR1", "T3_cor_above_cogs_consensus"]:
    m = tests.test == tname
    var = "cor_surprise_pct" if tname.startswith("T3") else "cost_surprise_pct"
    ok = all(float(tests[m & (tests.window == w) & (tests.variable == var)].p_one_sided.iloc[0]) <= 0.10 and
             (tname != "T2_persistence_AR1" or float(tests[m & (tests.window == w) & (tests.variable == var)].stat.iloc[0]) > 0) for w in ["W1", "W2"])
    tests.loc[m & tests.window.isin(["W1", "W2"]), "verdict_both_windows"] = "PASS" if ok else "FAIL"
tests.to_csv(OUT / "41_tests.csv", index=False)

# ----------------------------------------------------------------------------------------------------------------------
# 3. Where our cost excess sits: cost of revenue vs LSEG's COGS consensus, the rest is opex
# ----------------------------------------------------------------------------------------------------------------------
att = []
lbann = lba[lba.scenario == "base"].set_index("period")
for per in ["3Q26", "4Q26", "FY26", "FY27"]:
    r = cur.loc[per]; st_costs = float(r.revenue_mean - r.ebitda_mean); st_cogs = float(r.cogs_mean)
    if per.startswith("FY"):
        o = lbann.loc[per]; da = float(o.da); our_costs = float(o.revenue - o.adj_ebitda); our_cor = float(o.cor_cash)
    else:
        o = lbq.loc[per]; da = float(o.da); our_costs = float(o.revenue - o.adj_ebitda); our_cor = float(o.cor_cash)
    att.append(dict(period=per, street_costs=st_costs, our_costs=our_costs, d_costs=our_costs - st_costs, street_cogs=st_cogs, street_cogs_obs_date=r.cogs_obs_date,
                    our_cor=our_cor, d_cor=our_cor - st_cogs, street_opex=st_costs - st_cogs, our_opex=our_costs - our_cor, d_opex=(our_costs - our_cor) - (st_costs - st_cogs),
                    our_hosting=float(lbq.loc[per, "cor_hosting"]) if per in lbq.index else float(lbq.loc[Q27, "cor_hosting"].sum()) if per == "FY27" else np.nan,
                    street_cogs_growth_pct=np.nan))
att = pd.DataFrame(att)
att.loc[att.period == "FY27", "street_cogs_growth_pct"] = (float(cur.loc["FY27", "cogs_mean"]) / float(cur.loc["FY26", "cogs_mean"]) - 1) * 100
att.loc[att.period == "FY27", "our_cor_growth_pct"] = (float(lbann.loc["FY27", "cor_cash"]) / float(lbann.loc["FY26", "cor_cash"]) - 1) * 100
att.to_csv(OUT / "41_cogs_attribution.csv", index=False)

# ----------------------------------------------------------------------------------------------------------------------
# 4. FY27 cost-growth scenarios on the Street's own FY26 cost base; revenue at the Street and at official v2
# ----------------------------------------------------------------------------------------------------------------------
A = ann.set_index(["period", "source"])
st26 = A.loc[("FY26", "street_annual")]; st27 = A.loc[("FY27", "street_annual")]
g_hist = float(A.loc[[("FY24", "actual"), ("FY25", "actual")], "cost_growth_pct"].mean())
recent_surprise = float(tests[(tests.test == "count_costs_above_street") & (tests.window == "W1") & (tests.variable == "cost_surprise_pct")].recent_mean.iloc[0])
shares_fy27 = float(lbq.loc[Q27, "diluted_shares_m"].mean())
scen = {
    "Street (implied)": float(st27.cost_growth_pct),
    "Street + recent cost surprise (+%.2f%% of costs)" % recent_surprise: (1 + st27.cost_growth_pct / 100) * (1 + recent_surprise / 100) * 100 - 100,
    "Line build base": float(A.loc[("FY27", "line_build_base"), "cost_growth_pct"]),
    "FY24-25 actual average": g_hist,
    "FY26E Street (FY25A -> FY26 Street)": float(st26.cost_growth_pct),
}
srows = []
for name, g in scen.items():
    costs = float(st26.ebitda_costs) * (1 + g / 100)
    for rev_name, rev in [("Street revenue", float(st27.revenue)), ("official v2 revenue", float(A.loc[("FY27", "official_v2"), "revenue"]))]:
        e = rev - costs; d = e - float(st27.adj_ebitda)
        srows.append(dict(cost_path=name, cost_growth_pct=g, revenue_basis=rev_name, revenue=rev, ebitda_costs=costs, adj_ebitda=e, margin_pct=e / rev * 100,
                          d_ebitda_vs_street=d, d_ebitda_vs_street_pct=d / float(st27.adj_ebitda) * 100, d_margin_bp=(e / rev - float(st27.adj_ebitda) / float(st27.revenue)) * 1e4,
                          d_eps_vs_street=d * (1 - ETR_FY27 / 100) / shares_fy27))
pd.DataFrame(srows).to_csv(OUT / "41_fy27_scenarios.csv", index=False)

# 2H26: Street costs plus the recent surprise, at Street revenue and at official v2 revenue
h2 = []
for q in ["3Q26", "4Q26"]:
    r, e = street_q[q]; c = r - e
    for label, cc, rr in [("Street", c, r), ("Street costs + recent surprise, Street revenue", c * (1 + recent_surprise / 100), r),
                          ("Street costs + recent surprise, official v2 revenue", c * (1 + recent_surprise / 100), float(ov2.loc[q, "revenue"])),
                          ("line build base", float(lbq.loc[q, "revenue"] - lbq.loc[q, "adj_ebitda"]), float(lbq.loc[q, "revenue"])),
                          ("official v2", float(ov2.loc[q, "revenue"] - ov2.loc[q, "adj_ebitda"]), float(ov2.loc[q, "revenue"]))]:
        h2.append(dict(quarter=q, case=label, revenue=rr, ebitda_costs=cc, adj_ebitda=rr - cc, margin_pct=(rr - cc) / rr * 100, d_ebitda_vs_street=(rr - cc) - e))
pd.DataFrame(h2).to_csv(OUT / "41_2h26_scenarios.csv", index=False)

live = pd.DataFrame([dict(item="3Q26 Street EBITDA costs (LSEG 11 Sep: revenue - adj. EBITDA)", value=street_q["3Q26"][0] - street_q["3Q26"][1]),
                     dict(item="3Q26 line build base EBITDA costs", value=float(lbq.loc["3Q26", "revenue"] - lbq.loc["3Q26", "adj_ebitda"])),
                     dict(item="3Q26 Street costs + recent surprise", value=(street_q["3Q26"][0] - street_q["3Q26"][1]) * (1 + recent_surprise / 100)),
                     dict(item="AR(1) persistence forecast of the 3Q26 cost surprise, % of Street costs (W2 fit)", value=float(tests[(tests.test == "T2_persistence_AR1") & (tests.window == "W2") & (tests.variable == "cost_surprise_pct")].pred_next_print.iloc[0])),
                     dict(item="AR(1) persistence forecast of the 3Q26 cost surprise, % of Street costs (W1 fit)", value=float(tests[(tests.test == "T2_persistence_AR1") & (tests.window == "W1") & (tests.variable == "cost_surprise_pct")].pred_next_print.iloc[0])),
                     dict(item="call registered 22 Sep: actual 3Q26 costs above the Street (1 = above); scored 5 Nov 2026", value=1.0)])
live.to_csv(OUT / "41_live_3q26.csv", index=False)

pd.set_option("display.width", 250); pd.set_option("display.max_columns", 30)
print("=== annual ===\n", ann.round(2).to_string(index=False))
print("\n=== quarterly ===\n", pd.DataFrame(qrows).round(2).to_string(index=False))
print("\n=== cost surprise history ===\n", s[["print_quarter", "revenue_surprise", "ebitda_surprise", "cost_surprise", "cost_surprise_pct", "cost_surprise_exvar_pct", "street_cogs", "actual_cor", "cor_surprise", "opex_surprise"]].round(1).to_string(index=False))
print("\n=== tests ===\n", tests.round(4).to_string(index=False))
print("\n=== COGS attribution ===\n", att.round(1).to_string(index=False))
print("\n=== FY27 scenarios ===\n", pd.DataFrame(srows).round(2).to_string(index=False))
print("\n=== 2H26 ===\n", pd.DataFrame(h2).round(1).to_string(index=False))
print("\nwrote", OUT)
